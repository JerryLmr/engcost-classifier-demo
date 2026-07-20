from __future__ import annotations

import csv
import importlib.util
import json
import re
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

try:
    import numpy as np
    import pandas as pd
except ImportError:
    np = None
    pd = None


ROOT = Path(__file__).resolve().parents[3]


def load_script_module(name: str, relative_path: str):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative_path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


build_samples_script = load_script_module(
    "build_cost_item_samples", "estimator/scripts/build_cost_item_samples.py"
)
run_ingest_batch = load_script_module("run_ingest_batch", "estimator/scripts/run_ingest_batch.py")
merge_samples = load_script_module(
    "merge_cost_item_sample_batches", "estimator/scripts/merge_cost_item_sample_batches.py"
)

if np is not None and pd is not None:
    build_index = load_script_module(
        "build_cost_item_embedding_index", "estimator/scripts/build_cost_item_embedding_index.py"
    )
    query_estimate_llm = load_script_module(
        "query_cost_estimate_llm", "estimator/scripts/query_cost_estimate_llm.py"
    )
else:
    build_index = None
    query_estimate_llm = None


class MergeCostItemSampleBatchesTestCase(unittest.TestCase):
    SAMPLE_HEADERS = [
        "source_row_id",
        "file_name",
        "consultation_project_name",
        "sub_project_id",
        "seq",
        "project_name",
        "project_description",
        "unit",
        "quantity",
    ]

    def write_batch(self, input_dir: Path, batch_id: str, rows: list[dict[str, object]]) -> None:
        batch_dir = input_dir / batch_id
        batch_dir.mkdir(parents=True)
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "samples"
        sheet.append(self.SAMPLE_HEADERS)
        for row in rows:
            sheet.append([row.get(header) for header in self.SAMPLE_HEADERS])
        workbook.save(batch_dir / merge_samples.BATCH_SAMPLE_NAME)

    def sample_row(
        self,
        source_row_id: int,
        file_name: str,
        sub_project_id: str,
        project_name: str,
        *,
        consultation_project_name: str = "咨询项目",
        seq: int = 1,
    ) -> dict[str, object]:
        return {
            "source_row_id": source_row_id,
            "file_name": file_name,
            "consultation_project_name": consultation_project_name,
            "sub_project_id": sub_project_id,
            "seq": seq,
            "project_name": project_name,
            "project_description": f"{project_name}特征",
            "unit": "项",
            "quantity": 1,
        }

    def test_merge_keeps_latest_project_groups_then_latest_stable_rows(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)
            input_dir = tmp_path / "samples"
            output_path = tmp_path / "cost_item_samples_all.xlsx"
            report_path = merge_samples.dedup_report_path(output_path)

            old_blank_duplicate = self.sample_row(6, "", "空键子项目", "空键重复")
            latest_blank_duplicate = dict(old_blank_duplicate, source_row_id=16)
            same_batch_duplicate = self.sample_row(17, "同批次文件", "同批次子项目", "同批次重复")

            self.write_batch(
                input_dir,
                "20260618_001",
                [
                    self.sample_row(1, "楼幢项目", "7幢", "7幢旧数据"),
                    self.sample_row(2, "楼幢项目", "8幢", "8幢旧数据"),
                    self.sample_row(3, "楼幢项目", "9幢", "9幢旧数据一"),
                    self.sample_row(4, "楼幢项目", "9幢", "9幢旧数据二", seq=2),
                    self.sample_row(5, "另一个文件", "9幢", "不同文件保留"),
                    old_blank_duplicate,
                ],
            )
            self.write_batch(
                input_dir,
                "20260630_001",
                [
                    self.sample_row(10, "楼幢项目", "9幢", "9幢中间批次"),
                ],
            )
            self.write_batch(
                input_dir,
                "20260701_001",
                [
                    self.sample_row(15, "楼幢项目", "9幢", "9幢最新完整数据"),
                    latest_blank_duplicate,
                    same_batch_duplicate,
                    dict(same_batch_duplicate, source_row_id=18),
                ],
            )

            stats = merge_samples.merge_batches(input_dir, output_path, report_path)

            self.assertEqual(stats, (11, 6, 2, 3, 2))

            workbook = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
            sheet = workbook["samples"]
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            output_rows = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]
            workbook.close()

            self.assertEqual(
                headers,
                [*self.SAMPLE_HEADERS, "batch_id", "project_key", "stable_sample_id"],
            )
            output_names = [row["project_name"] for row in output_rows]
            self.assertEqual(
                output_names,
                [
                    "7幢旧数据",
                    "8幢旧数据",
                    "不同文件保留",
                    "9幢最新完整数据",
                    "空键重复",
                    "同批次重复",
                ],
            )
            self.assertNotIn("9幢旧数据一", output_names)
            self.assertNotIn("9幢旧数据二", output_names)
            self.assertNotIn("9幢中间批次", output_names)

            blank_row = next(row for row in output_rows if row["project_name"] == "空键重复")
            tie_row = next(row for row in output_rows if row["project_name"] == "同批次重复")
            self.assertEqual(blank_row["batch_id"], "20260701_001")
            self.assertEqual(blank_row["project_key"], "20260701_001::16")
            self.assertEqual(tie_row["project_key"], "20260701_001::17")
            for output_row in output_rows:
                source_row = {header: output_row[header] for header in self.SAMPLE_HEADERS}
                self.assertEqual(output_row["stable_sample_id"], merge_samples.stable_sample_id(source_row))
                self.assertEqual(
                    output_row["project_key"],
                    merge_samples.build_project_key(output_row["batch_id"], output_row["source_row_id"]),
                )

            with report_path.open("r", encoding="utf-8-sig", newline="") as report_file:
                report_rows = list(csv.DictReader(report_file))

            self.assertEqual(list(report_rows[0]), merge_samples.DEDUP_REPORT_HEADERS)
            group_reports = [row for row in report_rows if row["dedup_type"] == "project_group_replaced"]
            stable_reports = [row for row in report_rows if row["dedup_type"] == "stable_sample_duplicate"]
            self.assertEqual(len(group_reports), 2)
            self.assertEqual(
                {(row["duplicate_batch_id"], row["duplicate_row_count"]) for row in group_reports},
                {("20260618_001", "2"), ("20260630_001", "1")},
            )
            self.assertTrue(all(row["kept_batch_id"] == "20260701_001" for row in group_reports))
            self.assertTrue(all(row["dedup_key"] == "楼幢项目::9幢" for row in group_reports))
            self.assertTrue(all(row["stable_sample_id"] == "" for row in group_reports))
            self.assertTrue(all(row["cost_item_name"] == "" for row in group_reports))

            self.assertEqual(len(stable_reports), 2)
            self.assertTrue(all(row["duplicate_row_count"] == "1" for row in stable_reports))
            self.assertTrue(all(row["kept_batch_id"] == "20260701_001" for row in stable_reports))
            self.assertEqual(
                {row["duplicate_batch_id"] for row in stable_reports},
                {"20260618_001", "20260701_001"},
            )
            self.assertTrue(all(row["dedup_key"] == row["stable_sample_id"] for row in stable_reports))


@unittest.skipIf(np is None or pd is None, "cost item estimate dependencies are not installed")
class CostItemEstimateScriptTestCase(unittest.TestCase):
    def raw_sample_frame(self) -> pd.DataFrame:
        return pd.DataFrame(
            [
                {
                    "project_key": "batch-a::2",
                    "batch_id": "batch-a",
                    "source_row_id": 2,
                    "item_row_id": "2-1",
                    "工程名称": "屋面漏水维修工程",
                    "project_name_text": "屋面漏水维修",
                    "catalog_id": "CP-002-03",
                    "一级分类": "屋面",
                    "二级分类": "防水层",
                    "维修状态": "维修",
                    "标准对象": "共用部位",
                    "seq": 1,
                    "cost_item_name": "屋面卷材防水",
                    "project_description": "3.0mm SBS 沥青防水卷材",
                    "unit": "平方米",
                    "unit_normalized": "m²",
                    "quantity": 100,
                    "unit_price": 80,
                    "total_price": 8000,
                    "labor_unit_price": 20,
                    "machinery_unit_price": 5,
                    "consultation_time": "2026-03-01",
                    "location": "浙江省嘉兴市",
                },
                {
                    "project_key": "batch-a::2",
                    "batch_id": "batch-a",
                    "source_row_id": 2,
                    "item_row_id": "2-2",
                    "工程名称": "屋面漏水维修工程",
                    "project_name_text": "屋面漏水维修",
                    "catalog_id": "CP-002-03",
                    "一级分类": "屋面",
                    "二级分类": "防水层",
                    "维修状态": "维修",
                    "标准对象": "共用部位",
                    "seq": 2,
                    "cost_item_name": "防水层拆除",
                    "project_description": "拆除原屋面防水层",
                    "unit": "平方米",
                    "unit_normalized": "m²",
                    "quantity": 100,
                    "unit_price": 20,
                    "total_price": 2000,
                    "labor_unit_price": 12,
                    "machinery_unit_price": 1,
                    "consultation_time": "2026-03-01",
                    "location": "浙江省嘉兴市",
                },
                {
                    "project_key": "batch-a::3",
                    "batch_id": "batch-a",
                    "source_row_id": 3,
                    "item_row_id": "3-1",
                    "工程名称": "外墙维修工程",
                    "project_name_text": "外墙维修",
                    "catalog_id": "CP-003-01",
                    "一级分类": "外墙面",
                    "二级分类": "面层",
                    "维修状态": "维修",
                    "标准对象": "共用部位",
                    "seq": 1,
                    "cost_item_name": "外墙防水",
                    "project_description": "外墙渗漏处理",
                    "unit": "平方米",
                    "unit_normalized": "m²",
                    "quantity": 200,
                    "unit_price": 100,
                    "total_price": 20000,
                    "labor_unit_price": 30,
                    "machinery_unit_price": 8,
                    "consultation_time": "2024-01-01",
                    "location": "浙江省杭州市",
                },
                {
                    "project_key": "batch-a::4",
                    "batch_id": "batch-a",
                    "source_row_id": 4,
                    "item_row_id": "4-1",
                    "工程名称": "管道维修工程",
                    "project_name_text": "管道维修",
                    "catalog_id": "CF-015-04",
                    "一级分类": "给排水系统",
                    "二级分类": "管道",
                    "维修状态": "维修",
                    "标准对象": "共用设施设备",
                    "seq": 1,
                    "cost_item_name": "管道更换",
                    "project_description": "DN100 镀锌钢管更换",
                    "unit": "米",
                    "unit_normalized": "m",
                    "quantity": 30,
                    "unit_price": 100,
                    "total_price": 3000,
                    "labor_unit_price": "",
                    "machinery_unit_price": "",
                    "consultation_time": "2026-02-01",
                    "location": "浙江省嘉兴市",
                },
            ]
        )

    def prepared_samples(self) -> pd.DataFrame:
        samples = build_index.normalize_numeric_columns(
            build_index.ensure_optional_columns(build_index.ensure_project_key(self.raw_sample_frame()))
        )
        samples.insert(0, "sample_index", range(len(samples)))
        samples["project_package_id"] = samples["project_key"]
        samples["item_retrieval_text"] = samples.apply(build_index.build_item_retrieval_text, axis=1)
        samples["normalized_signature"] = samples.apply(build_index.build_normalized_signature, axis=1)
        return samples

    def test_normalize_embeddings_handles_zero_vector(self):
        embeddings = np.array([[3.0, 4.0], [0.0, 0.0]], dtype=np.float32)
        normalized = build_index.normalize_embeddings(embeddings)
        self.assertEqual(normalized.dtype, np.float32)
        self.assertAlmostEqual(float(np.linalg.norm(normalized[0])), 1.0)
        self.assertEqual(normalized[1].tolist(), [0.0, 0.0])

    def test_shared_normalize_unit_handles_square_and_cubic_units(self):
        self.assertEqual(build_samples_script.normalize_unit("m^2"), "m²")
        self.assertEqual(build_samples_script.normalize_unit("平方米"), "m²")
        self.assertEqual(build_samples_script.normalize_unit("平"), "m²")
        self.assertEqual(build_samples_script.normalize_unit("m^{3}"), "m³")
        self.assertEqual(build_samples_script.normalize_unit(" 台 "), "台")

    def test_item_retrieval_text_contains_only_item_description_and_unit(self):
        row = self.prepared_samples().iloc[0]

        self.assertEqual(
            row["item_retrieval_text"],
            "清单项：屋面卷材防水\n项目特征：3.0mm SBS 沥青防水卷材\n单位：m²",
        )
        self.assertNotIn("一级分类", row["item_retrieval_text"])
        self.assertNotIn("工程语义", row["item_retrieval_text"])
        self.assertEqual(row["normalized_signature"], "屋面卷材防水 | 3mm厚sbs防水卷材 | m²")

    def test_normalized_signature_normalizes_spacing_and_units(self):
        normalize = build_index.normalize_project_description

        self.assertEqual(normalize("外墙脚手架 高度13m以内"), normalize("外墙脚手架 高度 13m 以内"))
        self.assertEqual(normalize("外墙脚手架 高度20m以内"), normalize("外墙脚手架 高度 20m 以内"))
        self.assertNotEqual(normalize("外墙脚手架 高度13m以内"), normalize("外墙脚手架 高度20m以内"))
        self.assertEqual(build_index.normalize_unit("㎡"), build_index.normalize_unit("m2"))
        self.assertEqual(build_index.normalize_unit("平方米"), build_index.normalize_unit("m²"))

    def test_normalized_signature_normalizes_waterproof_thickness(self):
        normalize = build_index.normalize_project_description

        self.assertEqual(normalize("立面聚合物水泥防水涂料 ~1.2mm厚"), normalize("立面聚合物水泥防水涂料 ~1.2mm 厚"))
        self.assertEqual(normalize("平面聚氨酯防水涂料 ~1.5mm厚"), normalize("平面聚氨酯防水涂料~1.5mm 厚"))
        self.assertEqual(normalize("厚 1.5 mm 防水涂料"), normalize("1.5毫米厚防水涂料"))
        self.assertNotEqual(normalize("平面聚氨酯防水涂料 ~1.5mm厚"), normalize("平面聚氨酯防水涂料 ~1.2mm厚"))
        self.assertNotEqual(normalize("3mm SBS"), normalize("4mm SBS"))
        self.assertNotEqual(normalize("自粘卷材"), normalize("热熔卷材"))
        self.assertNotEqual(normalize("一层防水"), normalize("两层防水"))

    def test_normalized_signature_normalizes_sbs_waterproof_aliases(self):
        normalize = build_index.normalize_project_description
        expected = "3mm厚sbs防水卷材"

        equivalent_values = [
            "1.3.0mm厚弹性体改性沥青防水卷材",
            "1.3.0厚弹性体改性沥青防水卷材",
            "1.3.0mm弹性体改性沥青防水卷材",
            "3.0mm厚弹性体改性沥青防水卷材",
            "厚 3.0 mm 弹性体改性沥青防水卷材",
            "3.0mm弹性改性沥青防水卷材",
            "3.0mm SBS防水卷材",
            "3.0mm SBS改性沥青防水卷材",
            "3.0mm SBS弹性体改性沥青防水卷材",
            "卷材品种、规格、厚度：3.0mm SBS防水卷材",
        ]

        for value in equivalent_values:
            with self.subTest(value=value):
                self.assertEqual(normalize(value), expected)

    def test_normalized_signature_keeps_price_sensitive_waterproof_differences(self):
        normalize = build_index.normalize_project_description
        base = normalize("3.0mm SBS防水卷材")

        different_values = [
            "4.0mm SBS防水卷材",
            "3.0mm自粘SBS防水卷材",
            "3.0mm热熔SBS防水卷材",
            "3.0mm单层SBS防水卷材",
            "3.0mm双层SBS防水卷材",
            "3.0mm一道SBS防水卷材",
            "3.0mm两道SBS防水卷材",
            "3.0mm耐根穿刺SBS防水卷材",
            "3.0mm SBS防水卷材附加层",
            "3.0mm SBS防水卷材含基层处理",
            "拆除3.0mm SBS防水卷材",
            "新做3.0mm SBS防水卷材",
            "平面3.0mm SBS防水卷材",
            "立面3.0mm SBS防水卷材",
            "砂面3.0mm SBS防水卷材",
        ]

        for value in different_values:
            with self.subTest(value=value):
                self.assertNotEqual(normalize(value), base)

    def test_build_normalized_signature_normalizes_sbs_waterproof_aliases(self):
        row = pd.Series(
            {
                "cost_item_name": "屋面卷材防水",
                "project_description": "卷材品种、规格、厚度：1.3.0mm厚SBS防水卷材",
                "unit": "平方米",
                "unit_normalized": "m²",
            }
        )

        self.assertEqual(
            build_index.build_normalized_signature(row),
            "屋面卷材防水 | 3mm厚sbs防水卷材 | m²",
        )

    def test_normalized_signature_removes_layout_whitespace(self):
        normalize = build_index.normalize_project_description
        self.assertEqual(normalize("抹灰面铲除 抹灰面 只拆除面层时"), normalize("抹灰面铲除抹灰面只拆除面层时"))

    def test_normalized_signature_normalizes_thickness_order_without_merging_meaning(self):
        normalize = build_index.normalize_project_description

        self.assertEqual(normalize("厚 1.5（mm）聚氨酯防水涂料"), normalize("1.5mm厚聚氨酯防水涂料"))
        self.assertEqual(normalize("1.1.5mm厚聚氨酯防水涂料"), normalize("1.5mm聚氨酯防水涂料"))
        self.assertEqual(build_index.normalize_unit("m^{2}"), "m²")
        self.assertNotEqual(normalize("原有面层铲除及垃圾外运"), normalize("原面层拆除及垃圾清运"))
        self.assertNotEqual(normalize("屋面卷材防水 1.5mm厚"), normalize("墙面卷材防水 1.5mm厚"))

    def test_normalized_signature_removes_glued_list_numbers_but_keeps_decimals(self):
        normalize = build_index.normalize_project_description
        multiline = "1.原屋面基层清理\n2.刷1.2mm厚防水涂料"
        glued = "1.原屋面基层清理2.刷1.2mm防水涂料"

        self.assertEqual(normalize(multiline), normalize(glued))
        self.assertIn("1.2mm厚", normalize(glued))

    def test_normalized_signature_normalizes_confirmed_terms_only(self):
        normalize = build_index.normalize_project_description

        self.assertEqual(
            normalize("3.0厚弹性体改性沥青防水卷材"),
            normalize("3mmSBS改性沥青防水卷材"),
        )
        self.assertEqual(normalize("3.0mm自粘性防水卷材"), normalize("3mm自粘防水卷材"))
        self.assertEqual(normalize("单组份聚氨酯"), normalize("单组分聚氨酯"))
        self.assertNotEqual(normalize("原屋面防水层拆除"), normalize("原有屋面防水层拆除"))
        self.assertNotEqual(normalize("原屋面卷材铲除"), normalize("原屋面卷材拆除"))
        self.assertNotEqual(normalize("含垃圾外运"), normalize("含垃圾清运"))

    def test_normalized_signature_keeps_empty_description_slot_and_normalizes_unit(self):
        row = {"cost_item_name": "屋面保温修复", "project_description": "", "unit": "M2"}
        self.assertEqual(build_index.build_normalized_signature(row), "屋面保温修复 |  | m²")

    def test_project_packages_use_cost_item_names_summary_for_embedding(self):
        packages = build_index.build_project_packages(self.prepared_samples())

        roof = packages[packages["project_package_id"] == "batch-a::2"].iloc[0]
        self.assertEqual(roof["cost_item_names_summary"], "屋面卷材防水；防水层拆除")
        self.assertEqual(roof["item_summary"], "屋面卷材防水；防水层拆除")
        self.assertEqual(
            roof["package_text"],
            "工程名称：屋面漏水维修工程\n工程语义：屋面漏水维修\n包含清单项：屋面卷材防水；防水层拆除",
        )
        self.assertNotIn("CP-002-03", roof["package_text"])
        self.assertNotIn("3.0mm SBS", roof["package_text"])
        self.assertNotIn("单位", roof["package_text"])
        self.assertIn("catalog_summary", packages.columns)

    def test_build_index_meta_describes_new_retrieval_fields(self):
        meta = build_index.build_index_meta(Path("samples.xlsx"), "demo-model", 4, 3, 2)

        self.assertEqual(meta["files"]["project_package_embeddings"], "project_package_embeddings.npy")
        self.assertEqual(meta["files"]["item_embeddings"], "item_embeddings.npy")
        self.assertIn("cost_item_name、project_description、unit_normalized", meta["field_descriptions"]["item_retrieval_text"])
        self.assertIn("cost_item_names_summary", meta["field_descriptions"]["package_text"])

    def test_write_index_outputs_new_files_and_removes_legacy_files(self):
        samples = self.prepared_samples()
        packages = build_index.build_project_packages(samples)
        package_embeddings = np.ones((len(packages), 2), dtype=np.float32)
        item_embeddings = np.ones((len(samples), 2), dtype=np.float32)
        meta = build_index.build_index_meta(Path("samples.xlsx"), "demo-model", len(samples), len(packages), 2)

        with tempfile.TemporaryDirectory() as tmpdir:
            output_dir = Path(tmpdir) / "index"
            output_dir.mkdir()
            for name in build_index.LEGACY_OUTPUT_FILES:
                (output_dir / name).write_text("legacy", encoding="utf-8")

            build_index.write_index(samples, packages, package_embeddings, item_embeddings, output_dir, meta)

            self.assertTrue((output_dir / "samples.parquet").exists())
            self.assertTrue((output_dir / "project_packages.parquet").exists())
            self.assertTrue((output_dir / "project_package_embeddings.npy").exists())
            self.assertTrue((output_dir / "item_embeddings.npy").exists())
            for name in build_index.LEGACY_OUTPUT_FILES:
                self.assertFalse((output_dir / name).exists())

    def test_build_index_validate_output_dir_requires_overwrite_for_existing_directory(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output_dir = Path(tmpdir) / "index"
            output_dir.mkdir()

            with self.assertRaisesRegex(ValueError, "输出已存在，请加 --overwrite"):
                build_index.validate_output_dir(output_dir, overwrite=False)
            build_index.validate_output_dir(output_dir, overwrite=True)

    def test_query_parse_args_uses_new_top_package_and_item_options(self):
        with patch.object(sys, "argv", ["query_cost_estimate_llm.py", "--text", "屋面漏水"]):
            args = query_estimate_llm.parse_args()

        self.assertEqual(args.index_dir, "embeddings")
        self.assertEqual(args.top_packages, 20)
        self.assertEqual(args.top_items, 300)
        self.assertFalse(hasattr(args, "display_" + "selection_limit"))
        self.assertFalse(hasattr(args, "display_" + "exploration_limit"))
        self.assertEqual(args.package_weight_temperature, 0.1)
        self.assertFalse(hasattr(args, "family_final_score_limit"))
        self.assertFalse(hasattr(args, "family_item_score_limit"))
        self.assertEqual(args.llm_check_timeout, 3.0)
        self.assertFalse(hasattr(args, "top_k"))
        self.assertFalse(hasattr(args, "project_name_weight"))

    def test_query_main_checks_llm_service_before_run_query(self):
        call_order: list[str] = []

        def check_service(timeout_seconds: float):
            call_order.append(f"check:{timeout_seconds}")

        def run_query(**_kwargs):
            call_order.append("run_query")
            return object()

        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "query.xlsx"
            with patch.object(
                sys,
                "argv",
                [
                    "query_cost_estimate_llm.py",
                    "--text",
                    "屋面漏水",
                    "--output",
                    str(output_path),
                    "--llm-check-timeout",
                    "1.5",
                ],
            ), patch.object(query_estimate_llm, "check_lmstudio_service", side_effect=check_service) as check_mock, patch.object(
                query_estimate_llm,
                "run_query",
                side_effect=run_query,
            ) as run_mock, patch.object(query_estimate_llm, "print_terminal_summary"):
                exit_code = query_estimate_llm.main()

        self.assertEqual(exit_code, 0)
        self.assertEqual(call_order, ["check:1.5", "run_query"])
        check_mock.assert_called_once_with(timeout_seconds=1.5)
        run_mock.assert_called_once()

    def test_query_main_keeps_success_after_final_explanation_fallback(self):
        result = types.SimpleNamespace(success=False, error_message="invalid explanation")
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "query.xlsx"
            with patch.object(
                sys,
                "argv",
                ["query_cost_estimate_llm.py", "--text", "屋面漏水", "--output", str(output_path)],
            ), patch.object(query_estimate_llm, "check_lmstudio_service"), patch.object(
                query_estimate_llm, "run_query", return_value=result
            ), patch.object(query_estimate_llm, "print_terminal_summary"), patch("builtins.print") as print_mock:
                exit_code = query_estimate_llm.main()

        self.assertEqual(exit_code, 0)
        print_mock.assert_not_called()

    def test_query_main_returns_before_run_query_when_llm_service_unavailable(self):
        with patch.object(sys, "argv", ["query_cost_estimate_llm.py", "--text", "屋面漏水"]), patch.object(
            query_estimate_llm,
            "check_lmstudio_service",
            side_effect=RuntimeError("LMSTUDIO_BASE_URL=http://127.0.0.1:1234/v1"),
        ) as check_mock, patch.object(query_estimate_llm, "run_query") as run_mock, patch("builtins.print") as print_mock:
            exit_code = query_estimate_llm.main()

        self.assertEqual(exit_code, 1)
        check_mock.assert_called_once_with(timeout_seconds=3.0)
        run_mock.assert_not_called()
        print_mock.assert_called_once_with(
            "[ERROR] LLM 服务不可用，请先启动 LM Studio Server，并检查 "
            "LMSTUDIO_BASE_URL=http://127.0.0.1:1234/v1"
        )

    def test_query_load_index_reads_new_files_and_validates_shapes(self):
        samples = self.prepared_samples()
        packages = build_index.build_project_packages(samples)
        with tempfile.TemporaryDirectory() as tmpdir:
            index_dir = Path(tmpdir)
            samples.to_parquet(index_dir / "samples.parquet", index=False)
            packages.to_parquet(index_dir / "project_packages.parquet", index=False)
            np.save(index_dir / "project_package_embeddings.npy", np.ones((len(packages), 2), dtype=np.float32))
            np.save(index_dir / "item_embeddings.npy", np.zeros((len(samples), 2), dtype=np.float32))
            (index_dir / "index_meta.json").write_text('{"model": "demo-model"}', encoding="utf-8")

            loaded_samples, loaded_packages, package_embeddings, item_embeddings, meta = query_estimate_llm.load_index(index_dir)

        self.assertEqual(len(loaded_samples), 4)
        self.assertEqual(len(loaded_packages), 3)
        self.assertEqual(package_embeddings.shape, (3, 2))
        self.assertEqual(item_embeddings.shape, (4, 2))
        self.assertEqual(meta["model"], "demo-model")

    def test_query_rewrite_outputs_two_queries_and_fallbacks_empty_item_query(self):
        llm_result = {
            "project_package_query_text": "屋面漏水维修工程 屋面卷材防水",
            "item_query_text": "",
            "location": "浙江省嘉兴市",
            "start_date": "2025-07-14",
            "end_date": "2026-07-14",
            "extra_analysis": [{"raw_text": "500平", "value": 500, "unit": "m²"}],
            "extra_specs": ["3mm SBS"],
            "likely_catalog": {"SHOULD": "IGNORE"},
        }
        with patch.object(query_estimate_llm, "request_llm_json", return_value=llm_result):
            rewrite, trace = query_estimate_llm.query_rewrite_for_embedding(
                "屋面漏水", current_date=query_estimate_llm.date(2026, 7, 14)
            )

        self.assertTrue(rewrite.success)
        self.assertEqual(rewrite.project_package_query_text, "屋面漏水维修工程 屋面卷材防水")
        self.assertEqual(rewrite.item_query_text, "屋面漏水维修工程 屋面卷材防水")
        self.assertEqual((rewrite.location, rewrite.start_date, rewrite.end_date), ("浙江省嘉兴市", "2025-07-14", "2026-07-14"))
        self.assertIn("item_query_text 为空", rewrite.notes[0])
        self.assertEqual(trace["stage"], "query_rewrite_for_embedding")
        self.assertIn("当前日期：2026-07-14", trace["prompt"])

        with patch.object(query_estimate_llm, "request_llm_json", side_effect=query_estimate_llm.LLMServiceError("down")):
            fallback, trace = query_estimate_llm.query_rewrite_for_embedding("屋面漏水")

        self.assertFalse(fallback.success)
        self.assertEqual(fallback.project_package_query_text, "屋面漏水")
        self.assertEqual(fallback.item_query_text, "屋面漏水")
        self.assertEqual((fallback.location, fallback.start_date, fallback.end_date), ("", "", ""))
        self.assertEqual(trace["parsed_status"], "failed")

    def test_query_constraint_validation_accepts_only_standard_locations_and_strict_dates(self):
        for location in ["浙江省嘉兴市", "浙江省杭州市", "江苏省苏州市", "上海市", "北京市"]:
            with self.subTest(location=location):
                self.assertEqual(query_estimate_llm.validate_query_constraints(location, "", "")[:3], (location, "", ""))

        for location in ["嘉兴", "嘉兴市", "平湖市", "浙江省", "浙江省嘉兴市/上海市"]:
            with self.subTest(location=location):
                normalized, start, end, notes = query_estimate_llm.validate_query_constraints(location, "", "")
                self.assertEqual((normalized, start, end), ("", "", ""))
                self.assertTrue(notes)

        self.assertEqual(
            query_estimate_llm.validate_query_constraints("", "2025-07-14", "2026-07-14")[:3],
            ("", "2025-07-14", "2026-07-14"),
        )
        self.assertEqual(
            query_estimate_llm.validate_query_constraints("", "2025-13-01", "2025-02-30")[:3],
            ("", "", ""),
        )
        self.assertEqual(
            query_estimate_llm.validate_query_constraints("", "2026-01-01", "2025-01-01")[:3],
            ("", "", ""),
        )

    def test_constraint_mask_combines_location_and_strict_date_bounds(self):
        rows = pd.DataFrame([
            {"location": "浙江省嘉兴市", "consultation_time": "2025-07-14"},
            {"location": " 浙江省嘉兴市 ", "consultation_time": "2026-07-14"},
            {"location": "上海市", "consultation_time": "2026-01-01"},
            {"location": "浙江省嘉兴市", "consultation_time": "2025-02-30"},
        ])
        self.assertEqual(query_estimate_llm.build_constraint_mask(rows, "", "", "").tolist(), [True] * 4)
        self.assertEqual(
            query_estimate_llm.build_constraint_mask(rows, "浙江省嘉兴市", "2025-07-14", "2026-07-14").tolist(),
            [True, True, False, False],
        )
        parsed = query_estimate_llm.parse_consultation_dates(rows["consultation_time"])
        self.assertEqual(int(parsed.isna().sum()), 1)

    def test_constraint_filter_keeps_dataframe_and_embedding_rows_aligned(self):
        rows = pd.DataFrame({"value": ["a", "b", "c"]}, index=[10, 20, 30])
        embeddings = np.array([[1, 0], [2, 0], [3, 0]], dtype=np.float32)
        mask = pd.Series([True, False, True], index=rows.index)
        filtered_rows, filtered_embeddings = query_estimate_llm.filter_rows_and_embeddings(
            rows, embeddings, mask, "测试数据"
        )
        self.assertEqual(filtered_rows["value"].tolist(), ["a", "c"])
        self.assertEqual(filtered_embeddings.tolist(), [[1.0, 0.0], [3.0, 0.0]])
        with self.assertRaisesRegex(ValueError, "过滤前"):
            query_estimate_llm.filter_rows_and_embeddings(rows, embeddings[:2], mask, "测试数据")

    def test_empty_constrained_packages_fail_without_fallback_and_empty_direct_items_stay_empty(self):
        rewrite = query_estimate_llm.QueryRewrite(
            "屋面", "屋面工程", "屋面防水", "浙江省嘉兴市", "2025-01-01", "2025-12-31", [], True
        )
        with self.assertRaisesRegex(ValueError, "没有找到同时满足地域和时间约束"):
            query_estimate_llm.ensure_project_package_candidates(pd.DataFrame(), rewrite)

        direct = query_estimate_llm.score_direct_items(
            pd.DataFrame(columns=["sample_index"]), np.array([], dtype=np.float32), top_items=300
        )
        self.assertTrue(direct.empty)

    def test_package_evidence_weights_use_continuous_softmax(self):
        weights = query_estimate_llm.build_package_evidence_weights(
            ["p1", "p2", "p3"],
            {"p1": 0.9, "p2": 0.6, "p3": 0.7},
            temperature=0.1,
        )

        self.assertEqual(weights.columns.tolist(), query_estimate_llm.PACKAGE_EVIDENCE_WEIGHT_COLUMNS)
        self.assertAlmostEqual(float(weights["package_evidence_weight"].sum()), 1.0)
        p1 = weights[weights["project_package_id"].eq("p1")].iloc[0]
        p2 = weights[weights["project_package_id"].eq("p2")].iloc[0]
        self.assertGreater(float(p1["package_evidence_weight"]), float(p2["package_evidence_weight"]))
        with self.assertRaisesRegex(ValueError, "必须大于 0"):
            query_estimate_llm.build_package_evidence_weights(["p1"], {"p1": 0.9}, temperature=0)
        empty = query_estimate_llm.build_package_evidence_weights([], {}, temperature=0.1)
        self.assertTrue(empty.empty)

    def test_package_recall_dedupes_repeated_packages_before_top_k(self):
        packages = build_index.build_project_packages(self.prepared_samples())
        duplicate = packages.iloc[0].copy()
        duplicate["project_package_id"] = "batch-z::99"
        packages = pd.concat([packages, pd.DataFrame([duplicate])], ignore_index=True)
        embeddings = np.array([[0.9, 0.1], [0.2, 0.8], [0.1, 0.9], [1.0, 0.0]], dtype=np.float32)

        matched = query_estimate_llm.score_project_packages(
            packages,
            embeddings,
            np.array([1.0, 0.0], dtype=np.float32),
            top_packages=3,
        )

        self.assertEqual(len(matched), 3)
        self.assertEqual(matched["rank"].tolist(), [1, 2, 3])
        self.assertEqual(matched["工程名称"].tolist().count("屋面漏水维修工程"), 1)

    def test_build_retrieved_evidence_items_uses_base_similarities_without_composite_scores(self):
        samples = self.prepared_samples()
        packages = build_index.build_project_packages(samples)
        matched = packages[packages["project_package_id"].isin(["batch-a::2", "batch-a::3"])].copy()
        matched.insert(0, "package_query_similarity", [0.9, 0.6])
        matched.insert(0, "rank", [1, 2])
        direct = samples[samples["sample_index"].isin([0, 3])].copy()
        item_query_similarities = np.array([0.95, 0.2, 0.3, 0.7], dtype=np.float32)
        package_similarity_by_id = {"batch-a::2": 0.9, "batch-a::3": 0.6, "batch-a::4": 0.4}

        retrieved_evidence_items = query_estimate_llm.build_retrieved_evidence_items(
            samples,
            matched,
            direct,
            item_query_similarities,
            package_query_similarity_by_id=package_similarity_by_id,
        )

        self.assertEqual(sorted(retrieved_evidence_items["sample_index"].tolist()), [0, 1, 2, 3])
        roof = retrieved_evidence_items[retrieved_evidence_items["sample_index"] == 0].iloc[0]
        pipe = retrieved_evidence_items[retrieved_evidence_items["sample_index"] == 3].iloc[0]
        self.assertAlmostEqual(float(roof["package_query_similarity"]), 0.9)
        self.assertAlmostEqual(float(roof["item_query_similarity"]), 0.95)
        self.assertEqual(roof["source_ref"], "batch-a::2::2-1")
        self.assertEqual(pipe["package_query_similarity"], 0.4)
        self.assertEqual(pipe["direct_hit"], True)
        self.assertNotIn("final_score", retrieved_evidence_items.columns)
        self.assertNotIn("cooccur_score", retrieved_evidence_items.columns)

    def test_candidate_families_group_by_normalized_signature_only(self):
        candidates = self.prepared_samples().head(1).copy()
        same = candidates.iloc[0].copy()
        same["project_description"] = "3.0mm   SBS 沥青防水卷材"
        same["normalized_signature"] = build_index.build_normalized_signature(same)
        same["quantity"] = 120
        same["unit_price"] = 90
        same["total_price"] = 10800
        same["labor_unit_price"] = 22
        same["machinery_unit_price"] = 6
        same["project_key"] = "batch-a::5"
        same["project_package_id"] = "batch-a::5"
        same["item_row_id"] = "5-1"
        different = candidates.iloc[0].copy()
        different["project_description"] = "4.0mm SBS 沥青防水卷材"
        different["normalized_signature"] = build_index.build_normalized_signature(different)
        different["quantity"] = 300
        different["unit_price"] = 120
        different["total_price"] = 36000
        different["labor_unit_price"] = 30
        different["machinery_unit_price"] = 7
        different["project_key"] = "batch-a::6"
        different["project_package_id"] = "batch-a::6"
        different["item_row_id"] = "6-1"
        candidates = pd.concat([candidates, same.to_frame().T, different.to_frame().T], ignore_index=True)
        candidates["package_query_similarity"] = [0.8, 0.7, 0.6]
        candidates["package_rank"] = [1, 2, 3]
        candidates["item_query_similarity"] = [0.7, 0.9, 0.6]
        candidates["source_ref"] = ["batch-a::2::2-1", "batch-a::5::5-1", "batch-a::6::6-1"]

        with patch("builtins.print") as print_mock:
            families = query_estimate_llm.build_candidate_families(candidates)
        print_mock.assert_not_called()

        self.assertEqual(families.columns.tolist(), query_estimate_llm.CANDIDATE_FAMILY_COLUMNS)
        self.assertEqual(len(families), 2)
        roof3 = families[families["representative_project_description"].str.contains("3.0mm")].iloc[0]
        roof4 = families[families["representative_project_description"].str.contains("4.0mm")].iloc[0]
        self.assertEqual(roof3["normalized_signature"], build_index.build_normalized_signature(candidates.iloc[0]))
        self.assertEqual(roof3["本次召回样本数"], 2)
        self.assertEqual(roof3["本次召回工程包数"], 2)
        self.assertEqual(roof3["本次召回综合单价最低值"], 80.0)
        self.assertEqual(roof3["本次召回综合单价中位数"], 85.0)
        self.assertEqual(roof3["本次召回综合单价最高值"], 90.0)
        self.assertEqual(roof3["本次召回人工费单价中位数"], 21.0)
        self.assertEqual(roof3["本次召回机械费单价中位数"], 5.5)
        self.assertEqual(roof3["package_query_similarity最大值"], 0.8)
        self.assertEqual(roof3["item_query_similarity最大值"], 0.9)
        self.assertEqual(roof3["representative_project_description"], "3.0mm SBS 沥青防水卷材")
        self.assertEqual(roof3["family_id"], "F001")
        self.assertEqual(roof4["family_id"], "F002")
        self.assertEqual(roof4["本次召回综合单价最低值"], 120.0)

        evidence = query_estimate_llm.attach_family_ids_to_evidence_items(candidates, families)
        family_by_signature = dict(zip(families["normalized_signature"], families["family_id"], strict=False))
        self.assertEqual(len(evidence), len(candidates))
        self.assertIn("family_id", evidence.columns)
        self.assertIn("normalized_signature", evidence.columns)
        self.assertFalse(evidence["family_id"].map(query_estimate_llm.cell_text).eq("").any())
        for _index, row in evidence.iterrows():
            self.assertEqual(row["family_id"], family_by_signature[row["normalized_signature"]])

    def test_display_groups_group_by_display_name_and_unit_without_merging_prices(self):
        families = pd.DataFrame(
            [
                {
                    "family_id": "F001",
                    "normalized_signature": "sig-1",
                    "representative_cost_item_name": "屋面卷材防水",
                    "representative_project_description": "3.0mm SBS",
                    "unit": "平方米",
                    "unit_normalized": "m²",
                    "本次召回样本数": 2,
                    "本次召回工程包数": 2,
                    "本次召回综合单价中位数": 80,
                    "item_query_similarity最大值": 0.8,
                },
                {
                    "family_id": "F002",
                    "normalized_signature": "sig-2",
                    "representative_cost_item_name": "屋面卷材防水",
                    "representative_project_description": "4.0mm SBS",
                    "unit": "m²",
                    "unit_normalized": "m²",
                    "本次召回样本数": 3,
                    "本次召回工程包数": 3,
                    "本次召回综合单价中位数": 120,
                    "item_query_similarity最大值": 0.7,
                },
                {
                    "family_id": "F003",
                    "normalized_signature": "sig-3",
                    "representative_cost_item_name": "屋面卷材防水",
                    "representative_project_description": "含基层处理",
                    "unit": "m²",
                    "unit_normalized": "m²",
                    "本次召回样本数": 1,
                    "本次召回工程包数": 1,
                    "本次召回综合单价中位数": 140,
                    "item_query_similarity最大值": 0.6,
                },
                {
                    "family_id": "F004",
                    "normalized_signature": "sig-4",
                    "representative_cost_item_name": "垃圾外运",
                    "representative_project_description": "建筑垃圾外运",
                    "unit": "项",
                    "unit_normalized": "项",
                    "本次召回样本数": 1,
                    "本次召回工程包数": 1,
                    "本次召回综合单价中位数": 500,
                    "item_query_similarity最大值": 0.5,
                },
            ],
            columns=query_estimate_llm.CANDIDATE_FAMILY_COLUMNS,
        )
        evidence = pd.DataFrame(
            [
                {"family_id": "F001", "project_package_id": "p1"},
                {"family_id": "F002", "project_package_id": "p2"},
                {"family_id": "F003", "project_package_id": "p2"},
                {"family_id": "F004", "project_package_id": "p3"},
            ]
        )

        groups, group_families = query_estimate_llm.build_candidate_display_groups(families, evidence)

        self.assertEqual(len(groups), 2)
        roof = groups[groups["display_name"].eq("屋面卷材防水")].iloc[0]
        self.assertEqual(roof["family_count"], 3)
        self.assertEqual(roof["family_ids"], "F001,F002,F003")
        self.assertEqual(roof["retrieval_package_count"], 2)
        self.assertEqual(roof["direct_item_similarity_max"], 0.8)
        self.assertEqual(set(group_families[group_families["display_id"].eq(roof["display_id"])]["family_id"]), {"F001", "F002", "F003"})
        self.assertEqual(families[families["family_id"].eq("F001")].iloc[0]["本次召回综合单价中位数"], 80)
        self.assertEqual(families[families["family_id"].eq("F002")].iloc[0]["本次召回综合单价中位数"], 120)

    def test_display_groups_do_not_merge_different_units(self):
        families = pd.DataFrame(
            [
                {"family_id": "F001", "representative_cost_item_name": "垃圾外运", "representative_project_description": "按项", "unit": "项", "unit_normalized": "项", "本次召回样本数": 1, "本次召回工程包数": 1, "item_query_similarity最大值": 0.9},
                {"family_id": "F002", "representative_cost_item_name": "垃圾外运", "representative_project_description": "按方", "unit": "m³", "unit_normalized": "m³", "本次召回样本数": 1, "本次召回工程包数": 1, "item_query_similarity最大值": 0.8},
            ],
            columns=query_estimate_llm.CANDIDATE_FAMILY_COLUMNS,
        )
        groups, _families = query_estimate_llm.build_candidate_display_groups(families, pd.DataFrame([{"family_id": "F001", "project_package_id": "p1"}, {"family_id": "F002", "project_package_id": "p2"}]))

        self.assertEqual(len(groups), 2)
        self.assertEqual(set(groups["unit"]), {"项", "m³"})

    def test_display_source_package_count_uses_evidence_distinct_packages(self):
        families = pd.DataFrame(
            [
                {"family_id": "F001", "representative_cost_item_name": "屋面卷材防水", "representative_project_description": "3mm", "unit": "m²", "unit_normalized": "m²", "本次召回样本数": 2, "本次召回工程包数": 2, "item_query_similarity最大值": 0.9},
                {"family_id": "F002", "representative_cost_item_name": "屋面卷材防水", "representative_project_description": "4mm", "unit": "m²", "unit_normalized": "m²", "本次召回样本数": 2, "本次召回工程包数": 2, "item_query_similarity最大值": 0.8},
            ],
            columns=query_estimate_llm.CANDIDATE_FAMILY_COLUMNS,
        )
        evidence = pd.DataFrame(
            [
                {"family_id": "F001", "project_package_id": "p1"},
                {"family_id": "F001", "project_package_id": "p2"},
                {"family_id": "F002", "project_package_id": "p2"},
                {"family_id": "F002", "project_package_id": "p3"},
            ]
        )

        groups, _families = query_estimate_llm.build_candidate_display_groups(families, evidence)

        self.assertEqual(len(groups), 1)
        self.assertEqual(groups.iloc[0]["retrieval_package_count"], 3)

    def test_display_support_ratio_counts_each_package_once(self):
        groups = pd.DataFrame(
            [
                {"display_id": "D001", "display_key": "a|m²", "display_name": "A", "unit": "m²", "family_count": 2, "family_ids": "F001,F002", "top_family_examples": "[]"},
                {"display_id": "D002", "display_key": "b|m²", "display_name": "B", "unit": "m²", "family_count": 1, "family_ids": "F003", "top_family_examples": "[]"},
            ],
            columns=query_estimate_llm.CANDIDATE_DISPLAY_GROUP_COLUMNS,
        )
        display_families = pd.DataFrame(
            [
                {"display_id": "D001", "family_id": "F001"},
                {"display_id": "D001", "family_id": "F002"},
                {"display_id": "D002", "family_id": "F003"},
            ]
        )
        evidence = pd.DataFrame(
            [
                {"family_id": "F001", "project_package_id": "p1", "item_query_similarity": 0.8},
                {"family_id": "F002", "project_package_id": "p1", "item_query_similarity": 0.7},
                {"family_id": "F002", "project_package_id": "p2", "item_query_similarity": 0.9},
                {"family_id": "F003", "project_package_id": "p3", "item_query_similarity": 0.6},
            ]
        )
        weights = pd.DataFrame(
            [
                {"project_package_id": "p1", "package_query_similarity": 0.9, "package_evidence_weight": 0.5},
                {"project_package_id": "p2", "package_query_similarity": 0.8, "package_evidence_weight": 0.3},
                {"project_package_id": "p3", "package_query_similarity": 0.7, "package_evidence_weight": 0.2},
            ],
            columns=query_estimate_llm.PACKAGE_EVIDENCE_WEIGHT_COLUMNS,
        )

        output = query_estimate_llm.attach_display_support_ratios(groups, display_families, evidence, weights)
        d001 = output[output["display_id"].eq("D001")].iloc[0]

        self.assertAlmostEqual(float(d001["retrieval_package_support_ratio"]), 0.8)
        self.assertEqual(d001["retrieval_item_count"], 3)
        self.assertEqual(d001["retrieval_package_count"], 2)
        self.assertAlmostEqual(float(d001["direct_item_similarity_max"]), 0.9)

    def test_normalize_display_description_repairs_only_leading_ocr_numbering(self):
        normalize = query_estimate_llm.normalize_display_description

        self.assertEqual(normalize("1.3.0mmSBS防水卷材"), "3.0mmSBS防水卷材")
        self.assertEqual(normalize("1.4.0mmSBS防水卷材"), "4.0mmSBS防水卷材")
        self.assertEqual(normalize("1.1.5mm聚氨酯"), "1.5mm聚氨酯")
        self.assertEqual(normalize("1.5mm聚氨酯"), "1.5mm聚氨酯")

    def test_source_ref_recovers_from_batch_source_row_and_seq(self):
        rows = self.prepared_samples().head(1).copy()
        rows["project_key"] = ""
        rows["item_row_id"] = ""
        warnings: list[str] = []

        recovered = query_estimate_llm.attach_source_refs(rows, warnings)

        self.assertEqual(recovered.iloc[0]["project_key"], "batch-a::2")
        self.assertEqual(recovered.iloc[0]["item_row_id"], "2-1")
        self.assertEqual(recovered.iloc[0]["source_ref"], "batch-a::2::2-1")
        self.assertIn("source_ref_recovered_from_batch_source_row", warnings)
        self.assertIn("source_ref_recovered_from_source_row_seq", warnings)

    def test_matched_project_examples_use_top_three_complete_packages_and_source_order(self):
        matched = pd.DataFrame(
            [
                {"rank": 3, "project_package_id": "P3", "工程名称": "三"},
                {"rank": 1, "project_package_id": "P1", "工程名称": "一"},
                {"rank": 4, "project_package_id": "P4", "工程名称": "四"},
                {"rank": 2, "project_package_id": "P2", "工程名称": "二"},
            ]
        )
        samples = pd.DataFrame(
            [
                {"project_package_id": "P1", "item_row_id": "7-10", "seq": 10, "cost_item_name": "第十项"},
                {"project_package_id": "P1", "item_row_id": "7-2", "seq": 2, "cost_item_name": "第二项"},
                {"project_package_id": "P1", "item_row_id": "7-1", "seq": 1, "cost_item_name": "第一项"},
                {"project_package_id": "P2", "item_row_id": "x", "seq": 2, "cost_item_name": "P2-2"},
                {"project_package_id": "P2", "item_row_id": "y", "seq": 1, "cost_item_name": "P2-1"},
                {"project_package_id": "P3", "item_row_id": "1", "seq": 1, "cost_item_name": "P3-1"},
                {"project_package_id": "P4", "item_row_id": "1", "seq": 1, "cost_item_name": "P4-1"},
            ]
        )
        samples["stable_sample_id"] = [f"sid-{index}" for index in range(len(samples))]
        samples["source_ref"] = ["duplicate-ref" if index < 2 else f"ref-{index}" for index in range(len(samples))]

        examples = query_estimate_llm.build_matched_project_examples(matched, samples)

        self.assertEqual([example["project_package_id"] for example in examples], ["P1", "P2", "P3"])
        self.assertEqual([item["cost_item_name"] for item in examples[0]["items"]], ["第一项", "第二项", "第十项"])
        self.assertEqual([item["cost_item_name"] for item in examples[1]["items"]], ["P2-1", "P2-2"])
        self.assertEqual(sum(len(example["items"]) for example in examples), 6)
        self.assertTrue(all(item["stable_sample_id"] for example in examples for item in example["items"]))
        self.assertEqual(examples[0]["items"][0]["source_ref"], "ref-2")

    def test_matched_project_examples_generate_source_ref_without_lookup(self):
        matched = pd.DataFrame([{"rank": 1, "project_package_id": "P1"}])
        samples = pd.DataFrame([{
            "project_package_id": "P1", "stable_sample_id": "sid-1",
            "project_key": "batch::2", "item_row_id": "2-1", "cost_item_name": "清单项",
        }])

        examples = query_estimate_llm.build_matched_project_examples(matched, samples)

        self.assertEqual(examples[0]["items"][0]["source_ref"], "batch::2::2-1")

    def test_matched_examples_keep_full_project_when_grouping_only_requires_display_a(self):
        plan_items = pd.DataFrame([{"display_id": "D-A"}])
        display_groups = pd.DataFrame([
            {"display_id": "D-A", "display_name": "A", "unit": "m²", "family_count": 1},
            {"display_id": "D-B", "display_name": "B", "unit": "m²", "family_count": 1},
        ])
        display_families = pd.DataFrame([
            {"display_id": "D-A", "family_id": "F-A", "unit": "m²"},
            {"display_id": "D-B", "family_id": "F-B", "unit": "m²"},
        ])
        candidate_families = pd.DataFrame([
            {"family_id": "F-A", "representative_cost_item_name": "A", "unit": "m²", "本次召回样本数": 1},
            {"family_id": "F-B", "representative_cost_item_name": "B", "unit": "m²", "本次召回样本数": 1},
        ], columns=query_estimate_llm.CANDIDATE_FAMILY_COLUMNS)
        matched = pd.DataFrame([{"rank": 1, "project_package_id": "P1"}])
        samples = pd.DataFrame([
            {"project_package_id": "P1", "stable_sample_id": "sid-a", "seq": 1, "cost_item_name": "Display A 清单"},
            {"project_package_id": "P1", "stable_sample_id": "sid-b", "seq": 2, "cost_item_name": "Display B 清单"},
        ])

        required, filtered_groups, filtered_families = query_estimate_llm.filter_required_display_groups(
            plan_items, display_groups, display_families
        )
        grouped, *_rest = query_estimate_llm.generate_display_option_grouping(
            filtered_groups, filtered_families, candidate_families, []
        )
        examples = query_estimate_llm.build_matched_project_examples(matched, samples)

        self.assertEqual(required, ["D-A"])
        self.assertEqual(grouped["display_id"].tolist(), ["D-A"])
        self.assertEqual(
            [item["cost_item_name"] for item in examples[0]["items"]],
            ["Display A 清单", "Display B 清单"],
        )

    def test_selected_items_strictly_attach_family_and_display_ids(self):
        selected = pd.DataFrame([
            {"stable_sample_id": "sid-1", "item_position": 0},
            {"stable_sample_id": "sid-2", "item_position": 1},
        ])
        evidence = pd.DataFrame([
            {"stable_sample_id": "sid-1", "family_id": "F001"},
            {"stable_sample_id": "sid-2", "family_id": "F002"},
        ])
        display_families = pd.DataFrame([
            {"family_id": "F001", "display_id": "D002"},
            {"family_id": "F002", "display_id": "D001"},
        ])

        result = query_estimate_llm.attach_family_and_display_ids_to_selected_items(
            selected, evidence, display_families
        )

        self.assertEqual(result["family_id"].tolist(), ["F001", "F002"])
        self.assertEqual(result["display_id"].tolist(), ["D002", "D001"])

    def test_selected_items_mapping_rejects_missing_and_duplicate_mappings(self):
        selected = pd.DataFrame([{"stable_sample_id": "sid-1"}])
        with self.assertRaisesRegex(ValueError, "每行必须唯一映射"):
            query_estimate_llm.attach_family_and_display_ids_to_selected_items(
                selected,
                pd.DataFrame([{"stable_sample_id": "sid-2", "family_id": "F001"}]),
                pd.DataFrame([{"family_id": "F001", "display_id": "D001"}]),
            )
        with self.assertRaisesRegex(ValueError, "family_id 必须唯一映射"):
            query_estimate_llm.attach_family_and_display_ids_to_selected_items(
                selected,
                pd.DataFrame([{"stable_sample_id": "sid-1", "family_id": "F001"}]),
                pd.DataFrame([
                    {"family_id": "F001", "display_id": "D001"},
                    {"family_id": "F001", "display_id": "D002"},
                ]),
            )

    def test_filter_required_display_groups_preserves_required_and_family_order(self):
        plan_items = pd.DataFrame([
            {"display_id": "D002"}, {"display_id": "D001"}, {"display_id": "D002"},
        ])
        groups = pd.DataFrame([
            {"display_id": "D001"}, {"display_id": "D002"}, {"display_id": "D003"},
        ])
        families = pd.DataFrame([
            {"display_id": "D001", "family_id": "F001"},
            {"display_id": "D002", "family_id": "F002"},
            {"display_id": "D002", "family_id": "F003"},
            {"display_id": "D003", "family_id": "F004"},
        ])

        required, filtered_groups, filtered_families = query_estimate_llm.filter_required_display_groups(
            plan_items, groups, families
        )

        self.assertEqual(required, ["D002", "D001"])
        self.assertEqual(filtered_groups["display_id"].tolist(), ["D002", "D001"])
        self.assertEqual(filtered_families["family_id"].tolist(), ["F002", "F003", "F001"])

    def test_attach_original_practice_options_is_strict(self):
        plan_items = pd.DataFrame([{"family_id": "F001"}, {"family_id": "F002"}])
        displays = pd.DataFrame([{
            "display_id": "D001",
            "practice_options": [
                {"practice_option_id": "D001-O01", "family_ids": ["F001"]},
                {"practice_option_id": "D001-O02", "family_ids": ["F002"]},
            ],
        }])
        result = query_estimate_llm.attach_original_practice_options(plan_items, displays)
        self.assertEqual(result["practice_option_id"].tolist(), ["D001-O01", "D001-O02"])

        with self.assertRaisesRegex(ValueError, "唯一映射"):
            query_estimate_llm.attach_original_practice_options(
                plan_items.head(1),
                pd.DataFrame([{
                    "practice_options": [
                        {"practice_option_id": "O1", "family_ids": ["F001"]},
                        {"practice_option_id": "O2", "family_ids": ["F001"]},
                    ],
                }]),
            )

    def grouping_result(self, family_ids, groups, tags=None):
        tags = tags or {}
        return {
            "families": [
                {"family_id": family_id, "thickness": "", "material": "", "level": "", **tags.get(family_id, {})}
                for family_id in family_ids
            ],
            "groups": groups,
        }

    def test_display_option_grouping_prompt_includes_family_coverage_contract(self):
        display_groups = pd.DataFrame([{"display_id": "D001", "display_name": "屋面卷材防水", "unit": "m²", "family_count": 2}])
        display_families = pd.DataFrame(
            [
                {
                    "display_id": "D001",
                    "family_id": "F001",
                    "representative_cost_item_name": "屋面卷材防水",
                    "representative_project_description": "A" * 130,
                    "unit": "m²",
                    "本次召回样本数": 1,
                    "本次召回工程包数": 1,
                    "item_query_similarity最大值": 0.8,
                },
                {
                    "display_id": "D001",
                    "family_id": "F002",
                    "representative_cost_item_name": "屋面卷材防水",
                    "representative_project_description": "4mm SBS",
                    "unit": "m²",
                    "本次召回样本数": 2,
                    "本次召回工程包数": 2,
                    "item_query_similarity最大值": 0.9,
                },
            ]
        )
        candidate_families = pd.DataFrame(
            [
                {"family_id": "F001", "本次召回综合单价最低值": 80, "本次召回综合单价中位数": 90, "本次召回综合单价最高值": 100},
                {"family_id": "F002", "本次召回综合单价最低值": 100, "本次召回综合单价中位数": 110, "本次召回综合单价最高值": 120},
            ],
            columns=query_estimate_llm.CANDIDATE_FAMILY_COLUMNS,
        )

        prompt, record = query_estimate_llm.build_display_option_grouping_prompt(
            display_groups, display_families, candidate_families
        )

        self.assertEqual(set(record), {"display_name", "candidate_families"})
        self.assertEqual(
            set(record["candidate_families"][0]),
            {"family_id", "name", "spec", "unit"},
        )
        self.assertEqual(record["candidate_families"][1]["spec"], "A" * 120)
        self.assertNotIn("display_id", prompt)
        self.assertNotIn("display_results", prompt)
        self.assertNotIn("candidate_family_ids", prompt)
        self.assertNotIn("candidate_family_count", prompt)
        for removed_field in ["samples", "packages", "item_query_similarity", "unit_price_min", "unit_price_median", "unit_price_max"]:
            self.assertNotIn(f'"{removed_field}"', prompt)
        for expected_text in [
            "先逐个提取",
            "thickness",
            "material",
            "level",
            "顶层只允许 families 和 groups",
            "3mm 与 4mm 必须拆分",
            "2mm 水泥基渗透结晶与 2mm 聚氨酯必须拆分",
            "3mm 自粘卷材与 3mm SBS改性沥青卷材必须拆分",
            "1层、2层、5层、未注明层数必须分别拆分",
            "1.5mm单组份聚氨酯与1.5mm厚单组分聚氨酯可以合并",
            "1.2mm聚合物水泥基，含基层清理",
            "不要增加 is_generic",
        ]:
            self.assertIn(expected_text, prompt)
        for removed_text in ["raw_query", "project_package_query_text", "item_query_text", "selection_reason", "default_representative_family_id", "family_assignments", "group_reason"]:
            self.assertNotIn(removed_text, prompt)

    def test_display_option_grouping_groups_practice_options_without_legacy_fields(self):
        display_groups = pd.DataFrame([{"display_id": "D001", "display_name": "屋面卷材防水", "unit": "m²", "family_count": 3}])
        display_families = pd.DataFrame(
            [
                {"display_id": "D001", "family_id": "F004", "unit": "m²", "本次召回样本数": 2},
                {"display_id": "D001", "family_id": "F007", "unit": "m²", "本次召回样本数": 3},
                {"display_id": "D001", "family_id": "F010", "unit": "m²", "本次召回样本数": 1},
            ]
        )

        displays_with_options, meta = query_estimate_llm.parse_display_option_grouping_result(
            self.grouping_result(["F004", "F007", "F010"], [["F007", "F004"], ["F010"]]),
            display_groups,
            display_families,
            [],
        )

        self.assertEqual(displays_with_options.loc[0, "practice_options"][0]["practice_option_id"], "D001-O01")
        self.assertEqual(displays_with_options.loc[0, "practice_options"][0]["sample_count"], 5)
        self.assertEqual(displays_with_options.loc[0, "practice_options"][1]["practice_option_id"], "D001-O02")
        self.assertNotIn("practice_description", displays_with_options.loc[0, "practice_options"][0])
        self.assertNotIn("representative_family_id", displays_with_options.loc[0, "practice_options"][0])
        self.assertNotIn("selected_family_id", displays_with_options.columns)
        self.assertNotIn("default_practice_option_id", displays_with_options.columns)
        self.assertNotIn("default_practice", displays_with_options.columns)
        self.assertNotIn("family_relations", displays_with_options.columns)
        self.assertNotIn("selection_reason", displays_with_options.columns)
        self.assertNotIn("group_reason", displays_with_options.loc[0, "practice_options"][0])
        self.assertEqual(meta["display_ids"], ["D001"])
        self.assertNotIn("default_practice_option_ids", meta)
        self.assertEqual(meta["practice_option_count"], 2)
        self.assertEqual(meta["families_grouped_count"], 3)
        self.assertEqual(meta["option_count_by_display"], {"D001": 2})
        self.assertEqual(meta["max_options_per_display"], 2)

    def test_display_option_grouping_covers_each_of_four_families_once(self):
        display_groups = pd.DataFrame([{"display_id": "D001", "display_name": "屋面防水", "unit": "m²", "family_count": 4}])
        display_families = pd.DataFrame(
            [{"display_id": "D001", "family_id": family_id, "unit": "m²"} for family_id in ["F001", "F002", "F003", "F004"]]
        )
        displays_with_options, meta = query_estimate_llm.parse_display_option_grouping_result(
            self.grouping_result(
                ["F001", "F002", "F003", "F004"],
                [["F003", "F004"], ["F002", "F001"]],
            ),
            display_groups,
            display_families,
            [],
        )

        flattened = [family_id for option in displays_with_options.loc[0, "practice_options"] for family_id in option["family_ids"]]
        self.assertEqual(flattened, ["F003", "F004", "F002", "F001"])
        self.assertEqual(len(flattened), len(set(flattened)))
        self.assertEqual(meta["families_grouped_count"], 4)

    def test_display_option_grouping_auto_split_conflict_matrix_and_stable_order(self):
        cases = [
            (
                "different thickness",
                ["F001", "F002"], [["F001", "F002"]],
                {"F001": {"thickness": "3mm"}, "F002": {"thickness": "4mm"}},
                [["F001"], ["F002"]], ["thickness_conflict"],
            ),
            (
                "one thickness and blanks",
                ["F001", "F002", "F003"], [["F001", "F002", "F003"]],
                {"F001": {"thickness": "3mm"}},
                [["F001", "F002", "F003"]], [],
            ),
            (
                "different material",
                ["F003", "F005", "F022"], [["F003", "F005", "F022"]],
                {"F003": {"material": "水泥基渗透结晶"}, "F005": {"material": "水泥基渗透结晶"}, "F022": {"material": "聚氨酯"}},
                [["F003", "F005"], ["F022"]], ["material_conflict"],
            ),
            (
                "one material and blanks",
                ["F001", "F012"], [["F001", "F012"]],
                {"F012": {"material": "聚合物水泥基"}},
                [["F001", "F012"]], [],
            ),
            (
                "different levels including blank",
                ["F086", "F087", "F088", "F093", "F094"], [["F086", "F087", "F088", "F093", "F094"]],
                {"F086": {"level": "2层"}, "F087": {"level": "5层"}, "F088": {"level": "1层"}},
                [["F086"], ["F087"], ["F088"], ["F093", "F094"]], ["level_conflict"],
            ),
            (
                "self adhesive versus sbs",
                ["F008", "F011"], [["F008", "F011"]],
                {"F008": {"thickness": "3mm", "material": "自粘卷材"}, "F011": {"thickness": "3mm", "material": "SBS改性沥青"}},
                [["F008"], ["F011"]], ["material_conflict"],
            ),
            (
                "same normalized polyurethane",
                ["F014", "F043"], [["F014", "F043"]],
                {"F014": {"thickness": "1.5mm", "material": "聚氨酯"}, "F043": {"thickness": "1.5mm", "material": "聚氨酯"}},
                [["F014", "F043"]], [],
            ),
            (
                "same tags with cleaning detail",
                ["F032", "F048"], [["F032", "F048"]],
                {"F032": {"thickness": "1.2mm", "material": "聚合物水泥基"}, "F048": {"thickness": "1.2mm", "material": "聚合物水泥基"}},
                [["F032", "F048"]], [],
            ),
            (
                "only conflicting group splits",
                ["F010", "F011", "F025", "F026", "F027"], [["F010", "F011"], ["F025", "F026", "F027"]],
                {"F010": {"thickness": "2mm"}, "F011": {"thickness": "2mm"}, "F025": {"thickness": "4mm"}, "F026": {"thickness": "3mm"}, "F027": {"thickness": "3mm"}},
                [["F010", "F011"], ["F025"], ["F026", "F027"]], ["thickness_conflict"],
            ),
            (
                "multiple conflicts",
                ["F001", "F002"], [["F001", "F002"]],
                {"F001": {"thickness": "3mm", "material": "SBS改性沥青", "level": "1层"}, "F002": {"thickness": "4mm", "material": "聚氨酯", "level": ""}},
                [["F001"], ["F002"]], ["multiple_conflicts"],
            ),
        ]
        for name, family_ids, groups, tags, expected_groups, expected_reasons in cases:
            with self.subTest(name=name):
                displays = pd.DataFrame([{"display_id": "D001", "display_name": "测试", "unit": "m²", "family_count": len(family_ids)}])
                display_families = pd.DataFrame([
                    {"display_id": "D001", "family_id": family_id, "unit": "m²", "本次召回样本数": 1}
                    for family_id in family_ids
                ])
                grouped, meta = query_estimate_llm.parse_display_option_grouping_result(
                    self.grouping_result(family_ids, groups, tags), displays, display_families, []
                )
                options = grouped.loc[0, "practice_options"]
                self.assertEqual([option["family_ids"] for option in options], expected_groups)
                self.assertEqual([option["practice_option_id"] for option in options], [f"D001-O{i:02d}" for i in range(1, len(expected_groups) + 1)])
                details = meta["grouping_details_by_display"]["D001"]
                self.assertEqual(details["final_groups"], expected_groups)
                self.assertEqual(details["auto_split_reasons"], expected_reasons)
                self.assertEqual(details["auto_split_applied"], bool(expected_reasons))

    def test_display_option_grouping_strict_validation_rejects_invalid_options(self):
        display_groups = pd.DataFrame([{"display_id": "D001", "display_name": "屋面防水", "unit": "m²", "family_count": 2}])
        display_families = pd.DataFrame(
            [
                {"display_id": "D001", "family_id": "F001", "unit": "m²"},
                {"display_id": "D001", "family_id": "F002", "unit": "m²"},
            ]
        )
        valid_result = self.grouping_result(["F001", "F002"], [["F001", "F002"]])

        def assert_invalid(result, message):
            with self.assertRaisesRegex(ValueError, message):
                query_estimate_llm.parse_display_option_grouping_result(
                    result,
                    display_groups,
                    display_families,
                    [],
                )

        assert_invalid({**valid_result, "display_id": "D999"}, "顶层只允许包含 families 和 groups")
        assert_invalid(
            {**valid_result, "groups": [["F001"]]},
            "遗漏 candidate family",
        )
        assert_invalid(
            {**valid_result, "groups": [["F001", "F001", "F002"]]},
            "family 重复",
        )
        assert_invalid(
            {**valid_result, "groups": [["F001", "BAD"]]},
            "不属于当前 display",
        )
        for legacy_field in ["default_representative_family_id", "selection_reason", "family_assignments"]:
            assert_invalid(
                {**valid_result, legacy_field: "legacy"},
                "顶层只允许包含 families 和 groups",
            )
        assert_invalid({**valid_result, "groups": []}, "groups 必须为非空")
        assert_invalid({**valid_result, "groups": [[]]}, "group 必须为非空")
        assert_invalid(
            {**valid_result, "groups": [["F001", "F002"], ["F002", "F001"]]},
            "完全相同的 family_ids 分组",
        )
        malformed_family = dict(valid_result["families"][0])
        malformed_family.pop("level")
        assert_invalid(
            {**valid_result, "families": [malformed_family, valid_result["families"][1]]},
            "family_id、thickness、material、level",
        )
        assert_invalid(
            {**valid_result, "families": [valid_result["families"][0], valid_result["families"][0]]},
            "为空或重复",
        )

    def test_display_option_grouping_rejects_unit_mismatch_inside_option(self):
        display_groups = pd.DataFrame([{"display_id": "D001", "display_name": "屋面防水", "unit": "m²", "family_count": 2}])
        display_families = pd.DataFrame(
            [
                {"display_id": "D001", "family_id": "F001", "unit": "m²"},
                {"display_id": "D001", "family_id": "F002", "unit": "项"},
            ]
        )

        with self.assertRaisesRegex(ValueError, "单位必须一致"):
            query_estimate_llm.parse_display_option_grouping_result(
                self.grouping_result(["F001", "F002"], [["F001", "F002"]]),
                display_groups,
                display_families,
                [],
            )

    def test_generate_display_option_grouping_calls_each_multi_family_display_independently(self):
        display_groups = pd.DataFrame(
            [
                {"display_id": "D001", "display_name": "屋面防水", "unit": "m²", "family_count": 1},
                {"display_id": "D002", "display_name": "墙面防水", "unit": "m²", "family_count": 2},
            ]
        )
        display_families = pd.DataFrame(
            [
                {"display_id": "D001", "family_id": "F001", "unit": "m²", "本次召回样本数": 1, "本次召回工程包数": 1, "item_query_similarity最大值": 0.9},
                {"display_id": "D002", "family_id": "F002", "unit": "m²", "本次召回样本数": 2, "本次召回工程包数": 2, "item_query_similarity最大值": 0.8},
                {"display_id": "D002", "family_id": "F003", "unit": "m²", "本次召回样本数": 3, "本次召回工程包数": 3, "item_query_similarity最大值": 0.7},
            ]
        )
        candidate_families = pd.DataFrame(
            [
                {"family_id": "F001", "representative_cost_item_name": "屋面防水", "representative_project_description": "3mm SBS", "unit": "m²", "unit_normalized": "m²", "本次召回样本数": 1},
                {"family_id": "F002", "representative_cost_item_name": "墙面防水", "representative_project_description": "1.5mm涂膜", "unit": "m²", "unit_normalized": "m²", "本次召回样本数": 2},
                {"family_id": "F003", "representative_cost_item_name": "墙面防水", "representative_project_description": "2.0mm涂膜", "unit": "m²", "unit_normalized": "m²", "本次召回样本数": 3},
            ],
            columns=query_estimate_llm.CANDIDATE_FAMILY_COLUMNS,
        )
        response = types.SimpleNamespace(
            content=self.grouping_result(
                ["F002", "F003"], [["F002", "F003"]],
                {"F002": {"thickness": "1.5mm"}, "F003": {"thickness": "2mm"}},
            ),
            usage={},
            raw_content="{}",
        )

        with patch.object(query_estimate_llm, "request_llm_json_with_usage", return_value=response) as llm_mock:
            (
                displays_with_options, _success, _fallback, _error, _prompt,
                _summary_trace, meta, trace_frame, per_display_traces,
            ) = query_estimate_llm.generate_display_option_grouping(
                display_groups, display_families, candidate_families, []
            )

        llm_mock.assert_called_once()
        self.assertEqual(displays_with_options["display_id"].tolist(), ["D001", "D002"])
        self.assertEqual(displays_with_options.iloc[0]["practice_options"][0]["practice_option_id"], "D001-O01")
        self.assertEqual(displays_with_options.iloc[1]["practice_options"][1]["practice_option_id"], "D002-O02")
        self.assertEqual(meta["programmatic_single_family_display_count"], 1)
        self.assertEqual(meta["llm_display_count"], 1)
        self.assertEqual(set(trace_frame["display_id"]), {"D001", "D002"})
        self.assertEqual(len(per_display_traces), 1)
        self.assertIn("D002", per_display_traces[0]["purpose"])
        self.assertEqual(json.loads(per_display_traces[0]["input_summary"]), {
            "display_id": "D002", "display_name": "墙面防水", "candidate_family_count": 2,
        })
        for field in [
            "display_id", "display_name", "candidate_family_count", "raw_response",
            "parsed_family_tags", "original_groups", "final_groups", "auto_split_applied",
            "auto_split_reasons", "status", "error", "prompt_tokens", "completion_tokens", "total_tokens",
        ]:
            self.assertIn(field, per_display_traces[0])
        self.assertEqual(per_display_traces[0]["status"], "success")
        self.assertTrue(per_display_traces[0]["auto_split_applied"])
        self.assertEqual(per_display_traces[0]["original_groups"], [["F002", "F003"]])
        self.assertEqual(per_display_traces[0]["final_groups"], [["F002"], ["F003"]])
        self.assertEqual(per_display_traces[0]["auto_split_reasons"], ["thickness_conflict"])

    def test_display_option_grouping_fallback_is_isolated_per_display(self):
        displays = pd.DataFrame([
            {"display_id": "D001", "display_name": "A", "unit": "m²", "family_count": 2},
            {"display_id": "D002", "display_name": "B", "unit": "m²", "family_count": 2},
        ])
        display_families = pd.DataFrame([
            {"display_id": "D001", "family_id": "F001", "unit": "m²", "本次召回样本数": 1, "本次召回工程包数": 1, "item_query_similarity最大值": .9},
            {"display_id": "D001", "family_id": "F002", "unit": "m²", "本次召回样本数": 1, "本次召回工程包数": 1, "item_query_similarity最大值": .8},
            {"display_id": "D002", "family_id": "F003", "unit": "m²", "本次召回样本数": 1, "本次召回工程包数": 1, "item_query_similarity最大值": .9},
            {"display_id": "D002", "family_id": "F004", "unit": "m²", "本次召回样本数": 1, "本次召回工程包数": 1, "item_query_similarity最大值": .8},
        ])
        families = pd.DataFrame([
            {"family_id": family_id, "representative_cost_item_name": family_id, "unit": "m²", "本次召回样本数": 1}
            for family_id in ["F001", "F002", "F003", "F004"]
        ], columns=query_estimate_llm.CANDIDATE_FAMILY_COLUMNS)
        responses = [
            types.SimpleNamespace(content=self.grouping_result(["F001", "F002"], [["F001", "F999"]]), usage={"prompt_tokens": 11, "completion_tokens": 2, "total_tokens": 13}, raw_content="bad"),
            types.SimpleNamespace(content=self.grouping_result(["F003", "F004"], [["F003", "F004"]]), usage={"prompt_tokens": 21, "completion_tokens": 3, "total_tokens": 24}, raw_content="good"),
        ]
        warnings = []
        with patch.object(query_estimate_llm, "request_llm_json_with_usage", side_effect=responses) as llm_mock:
            (
                grouped, success, fallback, error, _prompt, _summary_trace,
                _meta, _trace_frame, per_display_traces,
            ) = query_estimate_llm.generate_display_option_grouping(
                displays, display_families, families, warnings
            )
        self.assertEqual(llm_mock.call_count, 2)
        first_input = llm_mock.call_args_list[0].args[0].split("【输入数据】", 1)[1]
        second_input = llm_mock.call_args_list[1].args[0].split("【输入数据】", 1)[1]
        self.assertIn("F001", first_input)
        self.assertIn("F002", first_input)
        self.assertNotIn("F003", first_input)
        self.assertNotIn("F004", first_input)
        self.assertIn("F003", second_input)
        self.assertIn("F004", second_input)
        self.assertNotIn("F001", second_input)
        self.assertNotIn("F002", second_input)
        options_by_display = {row["display_id"]: row["practice_options"] for _index, row in grouped.iterrows()}
        self.assertEqual(len(options_by_display["D001"]), 2)
        self.assertEqual(len(options_by_display["D002"]), 1)
        self.assertFalse(success)
        self.assertTrue(fallback)
        self.assertIn("D001", error)
        self.assertNotIn("D002:", error)
        self.assertIn("display_option_grouping_fallback_single_family_options:D001", warnings)
        self.assertEqual([trace["parsed_status"] for trace in per_display_traces], ["failed", "success"])
        self.assertIn("fallback_single_family_options", per_display_traces[0]["error_message"])
        self.assertEqual(per_display_traces[0]["raw_response"], "bad")
        self.assertEqual(per_display_traces[1]["raw_response"], "good")
        self.assertEqual(per_display_traces[1]["error_message"], "")
        self.assertEqual(per_display_traces[0]["total_tokens"], 13)
        self.assertEqual(per_display_traces[1]["total_tokens"], 24)

    def displays_with_options_fixture(self):
        return pd.DataFrame(
            [
                {
                    "display_id": "D001",
                    "display_name": "屋面防水",
                    "unit": "m²",
                    "practice_options": [
                        {
                            "practice_option_id": "D001-O01",
                            "representative_family_id": "F001",
                            "practice_description": "3mm SBS卷材防水",
                            "sample_count": 3,
                            "family_ids": ["F001", "F002"],
                        },
                        {
                            "practice_option_id": "D001-O02",
                            "representative_family_id": "F003",
                            "practice_description": "聚氨酯涂膜防水",
                            "sample_count": 1,
                            "family_ids": ["F003"],
                        },
                    ],
                },
                {
                    "display_id": "D002",
                    "display_name": "防水层拆除",
                    "unit": "m²",
                    "practice_options": [
                        {
                            "practice_option_id": "D002-O01",
                            "representative_family_id": "F004",
                            "practice_description": "拆除原防水层",
                            "sample_count": 2,
                            "family_ids": ["F004"],
                        }
                    ],
                },
            ]
        )

    def test_display_option_grouping_trace_marks_practice_options_without_default_fields(self):
        displays_with_options = self.displays_with_options_fixture().iloc[[0]].copy()
        display_families = pd.DataFrame(
            [
                {"display_id": "D001", "display_name": "屋面防水", "family_id": family_id}
                for family_id in ["F001", "F002", "F003"]
            ]
        )
        families = pd.DataFrame(
            [
                {"family_id": "F001", "unit": "m²", "unit_normalized": "m²"},
                {"family_id": "F002", "unit": "m²", "unit_normalized": "m²"},
                {"family_id": "F003", "unit": "m²", "unit_normalized": "m²"},
            ],
            columns=query_estimate_llm.CANDIDATE_FAMILY_COLUMNS,
        )

        trace = query_estimate_llm.build_display_option_grouping_trace_frame(displays_with_options, display_families, families).set_index("family_id")

        self.assertIn("practice_option_id", query_estimate_llm.DISPLAY_OPTION_GROUPING_TRACE_COLUMNS)
        self.assertNotIn("is_current_default_option", query_estimate_llm.DISPLAY_OPTION_GROUPING_TRACE_COLUMNS)
        self.assertNotIn("included_in_current_price_scope", query_estimate_llm.DISPLAY_OPTION_GROUPING_TRACE_COLUMNS)
        self.assertNotIn("group_reason", query_estimate_llm.DISPLAY_OPTION_GROUPING_TRACE_COLUMNS)
        self.assertEqual(trace.index.tolist(), ["F001", "F002", "F003"])
        self.assertEqual(trace.loc["F001", "practice_option_id"], "D001-O01")
        self.assertEqual(trace.loc["F003", "practice_option_id"], "D001-O02")

    def test_representative_package_selection_uses_all_packages_top_five_and_stable_ties(self):
        packages = pd.DataFrame(
            [
                {"project_package_id": "P1", "package_query_similarity": 0.90, "item_count": 1},
                {"project_package_id": "P2", "package_query_similarity": 0.80, "item_count": 4},
                {"project_package_id": "P3", "package_query_similarity": 0.80, "item_count": 6},
                {"project_package_id": "P4", "package_query_similarity": 0.70, "item_count": 8},
                {"project_package_id": "P5", "package_query_similarity": 0.60, "item_count": 10},
                {"project_package_id": "P6", "package_query_similarity": 0.99, "item_count": 100},
            ]
        )
        selected, ranked = query_estimate_llm.select_representative_project_package(packages)

        self.assertAlmostEqual(ranked["item_count_average"].iloc[0], 21.5)
        self.assertEqual(ranked["project_package_id"].tolist()[:3], ["P6", "P1", "P2"])
        self.assertEqual(selected["project_package_id"], "P4")
        self.assertFalse(bool(ranked.loc[ranked["project_package_id"].eq("P5"), "is_top5_similarity"].iloc[0]))
        self.assertTrue(bool(ranked.loc[ranked["project_package_id"].eq("P4"), "is_selected_package"].iloc[0]))
        self.assertEqual(
            ranked.loc[ranked["package_query_similarity"].eq(0.80), "project_package_id"].tolist(),
            ["P2", "P3"],
        )

    def test_representative_package_selection_boundaries_and_distance_tie(self):
        for packages, expected in [
            (pd.DataFrame([{"project_package_id": "only", "package_query_similarity": 1, "item_count": 7}]), "only"),
            (pd.DataFrame([
                {"project_package_id": "first", "package_query_similarity": 0.9, "item_count": 4},
                {"project_package_id": "second", "package_query_similarity": 0.8, "item_count": 6},
            ]), "first"),
        ]:
            with self.subTest(expected=expected):
                selected, ranked = query_estimate_llm.select_representative_project_package(packages)
                self.assertEqual(selected["project_package_id"], expected)
                self.assertEqual(len(ranked), len(packages))
        with self.assertRaisesRegex(ValueError, "为空"):
            query_estimate_llm.select_representative_project_package(pd.DataFrame())

    def test_expand_selected_project_uses_all_items_original_order_and_positions(self):
        samples = pd.DataFrame([
            {"project_package_id": "P1", "stable_sample_id": "s3", "item_row_id": "7-3", "seq": 3},
            {"project_package_id": "OTHER", "stable_sample_id": "other", "item_row_id": "8-1", "seq": 1},
            {"project_package_id": "P1", "stable_sample_id": "s1", "item_row_id": "7-1", "seq": 1},
            {"project_package_id": "P1", "stable_sample_id": "s2", "item_row_id": "7-2", "seq": 2},
        ])
        selected = pd.Series({"project_package_id": "P1", "item_count": 3})
        output = query_estimate_llm.expand_selected_project_items(samples, selected)
        self.assertEqual(output["stable_sample_id"].tolist(), ["s1", "s2", "s3"])
        self.assertEqual(output["item_position"].tolist(), [0, 1, 2])
        with self.assertRaisesRegex(ValueError, "数量与 item_count 不一致"):
            query_estimate_llm.expand_selected_project_items(
                samples, pd.Series({"project_package_id": "P1", "item_count": 2})
            )

    def test_contiguous_range_validation(self):
        for payload, expected in [
            ({"start_item_position": 1, "end_item_position": 3}, (1, 3)),
            ({"start_item_position": 2, "end_item_position": 2}, (2, 2)),
            ({"start_item_position": 0, "end_item_position": 4}, (0, 4)),
        ]:
            self.assertEqual(query_estimate_llm.validate_contiguous_range(payload, 5), expected)
        invalid = [
            {"start_item_position": -1, "end_item_position": 2},
            {"start_item_position": 0, "end_item_position": 5},
            {"start_item_position": 3, "end_item_position": 2},
            {"start_item_position": "0", "end_item_position": 2},
            {"start_item_position": True, "end_item_position": 2},
            {"start_item_position": 0},
            [],
        ]
        for payload in invalid:
            with self.subTest(payload=payload), self.assertRaises(ValueError):
                query_estimate_llm.validate_contiguous_range(payload, 5)

    def test_range_prompt_uses_minimal_items_but_quantity_prompt_keeps_quantity_context(self):
        items = pd.DataFrame([{
            "item_position": 3,
            "cost_item_name": "屋面卷材防水",
            "project_description": "3mm SBS",
            "unit": "平方米",
            "unit_normalized": "m²",
            "quantity": 100,
        }])
        range_prompt = query_estimate_llm.build_contiguous_item_range_prompt("屋面维修", "历史工程", items)
        range_payload = json.loads(range_prompt.split("输入：\n", 1)[1])
        self.assertEqual(
            range_payload["items"],
            [{
                "item_position": 3,
                "cost_item_name": "屋面卷材防水",
                "project_description": "3mm SBS",
            }],
        )

        quantity_prompt = query_estimate_llm.build_quantity_determination_prompt("屋面维修", items)
        quantity_payload = json.loads(quantity_prompt.split("输入：\n", 1)[1])
        self.assertEqual(set(quantity_payload), {"user_query", "items"})
        self.assertEqual(quantity_payload["items"], [{
            "item_position": 3,
            "cost_item_name": "屋面卷材防水",
            "project_description": "3mm SBS",
            "unit": "平方米",
        }])

    def test_contiguous_range_failure_falls_back_to_full_project(self):
        items = pd.DataFrame([
            {"item_position": 0, "cost_item_name": "A"},
            {"item_position": 1, "cost_item_name": "B"},
            {"item_position": 2, "cost_item_name": "C"},
        ])
        failures = [
            RuntimeError("down"),
            types.SimpleNamespace(content={"bad": 1}, usage={}, raw_content="bad json"),
            types.SimpleNamespace(content={"start_item_position": 2, "end_item_position": 1}, usage={}, raw_content="invalid"),
        ]
        for failure in failures:
            with self.subTest(failure=repr(failure)), patch.object(
                query_estimate_llm,
                "request_llm_json_with_usage",
                side_effect=failure if isinstance(failure, Exception) else None,
                return_value=None if isinstance(failure, Exception) else failure,
            ):
                start, end, meta = query_estimate_llm.select_contiguous_item_range("需求", "工程", items)
                self.assertEqual((start, end), (0, 2))
                self.assertTrue(meta["fallback"])
                self.assertEqual(meta["range_selection_status"], "fallback_full_project")

    def quantity_rule_items(self):
        return pd.DataFrame([
            {"item_position": 0, "cost_item_name": "防水层拆除", "project_description": "", "unit": "m²", "quantity": 80},
            {"item_position": 4, "cost_item_name": "屋面卷材防水", "project_description": "4mm SBS", "unit": "m²", "quantity": 100},
            {"item_position": 6, "cost_item_name": "面层恢复", "project_description": "", "unit": "m²", "quantity": 60},
        ])

    def quantity_result(self, explicit_position: int | None = None):
        return {"items": [
            {
                "item_position": position,
                "quantity_source": "user_explicit" if position == explicit_position else "historical_median",
                "quantity": 500 if position == explicit_position else None,
                "explanation": f"清单{position}的工程量判断",
            }
            for position in [0, 4, 6]
        ]}

    def test_quantity_explicit_binding_does_not_copy_or_scale(self):
        items = self.quantity_rule_items()
        determinations = query_estimate_llm.validate_quantity_determination_result(
            self.quantity_result(explicit_position=4), items
        )
        stats = {
            0: {"sample_count": 3, "minimum": 10, "median": 20, "maximum": 30},
            4: {},
            6: {"sample_count": 3, "minimum": 40, "median": 60, "maximum": 80},
        }
        quantities = query_estimate_llm.calculate_quantities(items, determinations, stats, {})
        self.assertEqual(quantities[0]["quantity"]["value"], 20)
        self.assertEqual(quantities[4]["quantity"]["value"], 500)
        self.assertEqual(quantities[6]["quantity"]["value"], 60)
        self.assertEqual(quantities[4]["quantity_source"], "user_explicit")

    def test_quantity_result_strict_validation(self):
        items = self.quantity_rule_items()
        valid = self.quantity_result(explicit_position=4)
        parsed = query_estimate_llm.validate_quantity_determination_result(valid, items)
        self.assertEqual(set(parsed), {0, 4, 6})
        invalid_results = [
            {**valid, "extra": 1},
            {"items": valid["items"][:-1]},
            {"items": [*valid["items"], valid["items"][0]]},
            {"items": [{**valid["items"][0], "item_position": True}, *valid["items"][1:]]},
            {"items": [{**valid["items"][0], "quantity_source": "unknown"}, *valid["items"][1:]]},
            {"items": [{**valid["items"][0], "quantity": 1}, *valid["items"][1:]]},
            {"items": [valid["items"][0], {**valid["items"][1], "quantity": 0}, valid["items"][2]]},
            {"items": [{**valid["items"][0], "explanation": ""}, *valid["items"][1:]]},
            {"items": [{**valid["items"][0], "extra": 1}, *valid["items"][1:]]},
        ]
        for result in invalid_results:
            with self.subTest(result=result), self.assertRaises(ValueError):
                query_estimate_llm.validate_quantity_determination_result(result, items)

    def test_quantity_queries_without_direct_item_quantity_remain_historical_median(self):
        items = self.quantity_rule_items()
        queries = [
            "屋面墙面漏水，面积大概500平",
            "消防报警主机故障，需要更换",
            "电梯钢丝绳更换",
            "电梯维修",
        ]
        for query in queries:
            with self.subTest(query=query):
                prompt = query_estimate_llm.build_quantity_determination_prompt(query, items)
                self.assertEqual(json.loads(prompt.split("输入：\n", 1)[1])["user_query"], query)
                parsed = query_estimate_llm.validate_quantity_determination_result(
                    self.quantity_result(), items
                )
                self.assertEqual(
                    {entry["quantity_source"] for entry in parsed.values()},
                    {"historical_median"},
                )

    def test_quantity_statistics_median_filter_units_and_deduplicate(self):
        samples = pd.DataFrame([
            {"stable_sample_id": "s1", "unit": "m²", "quantity": 100},
            {"stable_sample_id": "s2", "unit": "平方米", "quantity": 200},
            {"stable_sample_id": "s3", "unit": "m²", "quantity": 10000},
            {"stable_sample_id": "s3", "unit": "m²", "quantity": 10000},
            {"stable_sample_id": "s4", "unit": "m", "quantity": 50},
            {"stable_sample_id": "s5", "unit": "m²", "quantity": None},
            {"stable_sample_id": "s6", "unit": "m²", "quantity": 0},
            {"stable_sample_id": "s7", "unit": "m²", "quantity": "abc"},
        ])
        stats = query_estimate_llm.build_quantity_statistics(samples, "m²")
        self.assertEqual(stats["sample_count"], 3)
        self.assertEqual(stats["minimum"], 100)
        self.assertEqual(stats["median"], 200)
        self.assertEqual(stats["maximum"], 10000)
        simple_stats = query_estimate_llm.build_quantity_statistics(pd.DataFrame([
            {"stable_sample_id": "a", "unit": "m²", "quantity": 100},
            {"stable_sample_id": "b", "unit": "m²", "quantity": 200},
            {"stable_sample_id": "c", "unit": "m²", "quantity": 300},
        ]), "m²")
        self.assertEqual(simple_stats["sample_count"], 3)
        self.assertEqual(simple_stats["median"], 200)

    def test_quantity_best_sample_fallback_and_failure(self):
        items = self.quantity_rule_items().iloc[[0]].copy()
        items["stable_sample_id"] = ["best"]
        determinations = {0: {
            "quantity_source": "historical_median", "quantity": None, "explanation": "没有用户数量",
        }}
        warnings = []
        quantities = query_estimate_llm.calculate_quantities(
            items, determinations,
            {0: {"sample_count": 0, "minimum": None, "median": None, "maximum": None}},
            {"best": {"quantity": 88, "unit": "m²"}}, warnings,
        )
        self.assertEqual(quantities[0]["quantity"]["value"], 88)
        self.assertTrue(quantities[0]["quantity_fallback_used"])
        self.assertIn("quantity_median_unavailable:0", warnings)
        self.assertIn("quantity_best_sample_fallback:0", warnings)
        with self.assertRaisesRegex(ValueError, "无有效工程量"):
            query_estimate_llm.calculate_quantities(
                items, determinations, {0: {"median": None}}, {"best": {"quantity": None, "unit": "m²"}}
            )

    def test_quantity_llm_failure_uses_historical_median_and_traces_items(self):
        items = pd.DataFrame([{
            "item_position": 4, "stable_sample_id": "best", "cost_item_name": "清单A",
            "project_description": "", "unit": "m²", "display_id": "D1",
            "selected_option_id": "D1-O01", "original_option_id": "D1-O01",
            "original_family_id": "F1", "representative_family_id": "F1",
        }])
        lookup = {"best": {
            "project_package_id": "P1", "source_ref": "best-ref", "display_id": "D1",
            "practice_option_id": "D1-O01", "family_id": "F1", "quantity": 10, "unit": "m²",
        }}
        displays = pd.DataFrame([{
            "display_id": "D1", "unit": "m²",
            "practice_options": [{"practice_option_id": "D1-O01", "family_ids": ["F1"]}],
        }])
        families = pd.DataFrame([{
            "family_id": "F1", "normalized_signature": "sig", "unit": "m²", "unit_normalized": "m²",
        }])
        samples = pd.DataFrame([
            {"stable_sample_id": "s1", "normalized_signature": "sig", "unit": "m²", "quantity": 100},
            {"stable_sample_id": "s2", "normalized_signature": "sig", "unit": "m²", "quantity": 200},
            {"stable_sample_id": "s3", "normalized_signature": "sig", "unit": "m²", "quantity": 300},
        ])
        warnings = []
        with patch.object(query_estimate_llm, "request_llm_json_with_usage", side_effect=RuntimeError("down")):
            scenario, _prompt, trace = query_estimate_llm.generate_quantity_determination(
                "未提供数量", "P1", items, lookup, displays, families, samples, warnings
            )
        self.assertEqual(scenario.items[0].quantity["value"], 200)
        self.assertEqual(scenario.items[0].quantity_source, "historical_median")
        self.assertTrue(trace["fallback"])
        self.assertEqual(json.loads(trace["quantity_items"])[0]["quantity_median"], 200)
        self.assertIn("quantity_determination_fallback_historical_median", warnings)

    def test_build_scenario_from_plan_preserves_historical_practice_option(self):
        plan_items = pd.DataFrame([{"item_position": 5, "stable_sample_id": "sid"}])
        lookup = {"sid": {
            "project_package_id": "P1", "source_ref": "ref", "display_id": "D1", "practice_option_id": "D1-O2"
        }}
        scenario = query_estimate_llm.build_scenario_from_plan_items(
            "P1", plan_items, lookup, {5: {
                "quantity": {"type": "exact", "value": 1},
                "quantity_reason": "依据", "quantity_source": "user_explicit",
                "quantity_explanation": "依据", "quantity_sample_count": None,
                "quantity_minimum": None, "quantity_median": None, "quantity_maximum": None,
                "quantity_fallback_used": False, "quantity_fallback_reason": "",
            }}
        )
        self.assertEqual(scenario.items[0].practice_option_id, "D1-O2")
        self.assertEqual(scenario.items[0].selection_reason, "")

    def test_option_selection_single_option_passthrough_and_empty_description(self):
        plan = pd.DataFrame([{"item_position": 0, "stable_sample_id": "sid"}])
        lookup = {"sid": {"family_id": "F1", "display_id": "D1", "practice_option_id": "D1-O01"}}
        displays = pd.DataFrame([{"display_id": "D1", "practice_options": [{"practice_option_id": "D1-O01", "family_ids": ["F1"]}]}])
        families = pd.DataFrame([{
            "family_id": "F1", "representative_cost_item_name": "电话主机",
            "representative_project_description": "", "unit": "台", "normalized_signature": "电话主机||台",
            "本次召回样本数": 1,
        }])
        with patch.object(query_estimate_llm, "request_llm_json_with_usage") as llm_mock:
            selected, trace, llm_traces = query_estimate_llm.select_final_options(
                "消防主机", plan, lookup, displays, families, []
            )
        llm_mock.assert_not_called()
        self.assertEqual(selected.loc[0, "project_description"], "")
        self.assertEqual(selected.loc[0, "representative_family_id"], "F1")
        self.assertEqual(trace.loc[0, "candidate_option_count"], 1)
        self.assertEqual(llm_traces, [])

    def test_option_selection_replacement_uses_deterministic_real_family(self):
        plan = pd.DataFrame([{"item_position": 0, "stable_sample_id": "sid"}])
        lookup = {"sid": {"family_id": "F1", "display_id": "D1", "practice_option_id": "D1-O01"}}
        displays = pd.DataFrame([{"display_id": "D1", "practice_options": [
            {"practice_option_id": "D1-O01", "family_ids": ["F1"]},
            {"practice_option_id": "D1-O02", "family_ids": ["F2", "F3"]},
        ]}])
        families = pd.DataFrame([
            {"family_id": "F1", "representative_cost_item_name": "屋面防水", "representative_project_description": "4mm SBS", "unit": "m²", "normalized_signature": "roof|4mm|m²", "本次召回样本数": 9},
            {"family_id": "F2", "representative_cost_item_name": "屋面防水", "representative_project_description": "", "unit": "m²", "normalized_signature": "roof||m²", "本次召回样本数": 3},
            {"family_id": "F3", "representative_cost_item_name": "屋面防水", "representative_project_description": "3mm SBS", "unit": "m²", "normalized_signature": "roof|3mm|m²", "本次召回样本数": 3},
        ])
        response = types.SimpleNamespace(
            content={"decision": "explicit_match", "selected_option_id": "D1-O02"}, usage={}, raw_content="{}"
        )
        with patch.object(query_estimate_llm, "request_llm_json_with_usage", return_value=response) as llm_mock:
            selected, trace, _llm_traces = query_estimate_llm.select_final_options(
                "使用3mm SBS", plan, lookup, displays, families, []
            )
        prompt = llm_mock.call_args.args[0]
        self.assertNotIn('"family_id"', prompt)
        self.assertNotIn('"normalized_signature"', prompt)
        self.assertIn('"original_option_id": "D1-O01"', prompt)
        self.assertIn('"option_id": "D1-O01"', prompt)
        self.assertIn('"option_id": "D1-O02"', prompt)
        self.assertEqual(selected.loc[0, "representative_family_id"], "F2")
        self.assertTrue(trace.loc[0, "whether_replaced"])

    def test_option_support_counts_deduplicate_packages_across_families(self):
        displays = pd.DataFrame([{"display_id": "D1", "practice_options": [
            {"practice_option_id": "D1-O01", "family_ids": ["F1", "F2"]},
            {"practice_option_id": "D1-O02", "family_ids": ["F3"]},
        ]}])
        evidence = pd.DataFrame([
            {"family_id": "F1", "project_package_id": "P1"},
            {"family_id": "F2", "project_package_id": "P1"},
            {"family_id": "F2", "project_package_id": "P2"},
            {"family_id": "F3", "project_package_id": ""},
        ])
        supported = query_estimate_llm.attach_option_support_counts(displays, evidence)
        options = supported.loc[0, "practice_options"]
        self.assertEqual((options[0]["option_sample_count"], options[0]["option_package_count"]), (3, 2))
        self.assertEqual((options[1]["option_sample_count"], options[1]["option_package_count"]), (1, 0))

    def test_most_supported_option_uses_sample_package_original_and_stable_ties(self):
        choose = query_estimate_llm.choose_most_supported_option
        options = [
            {"practice_option_id": "O1", "option_sample_count": 2, "option_package_count": 1},
            {"practice_option_id": "O2", "option_sample_count": 3, "option_package_count": 1},
        ]
        self.assertEqual(choose(options, "O1"), ("O2", "no_explicit_match_selected_by_sample_count"))
        options[0]["option_sample_count"] = 3
        options[0]["option_package_count"] = 2
        self.assertEqual(choose(options, "O2"), ("O1", "no_explicit_match_selected_by_package_count"))
        options[1]["option_package_count"] = 2
        self.assertEqual(choose(options, "O2"), ("O2", "support_tie_original_option"))
        self.assertEqual(choose(options, "missing"), ("O1", "stable_option_id_tiebreak"))

    def test_no_explicit_match_and_llm_failure_select_by_support(self):
        plan = pd.DataFrame([{"item_position": 0, "stable_sample_id": "sid"}])
        lookup = {"sid": {"family_id": "F1", "display_id": "D1", "practice_option_id": "D1-O01"}}
        displays = pd.DataFrame([{"display_id": "D1", "practice_options": [
            {"practice_option_id": "D1-O01", "family_ids": ["F1"]},
            {"practice_option_id": "D1-O02", "family_ids": ["F2"]},
        ]}])
        families = pd.DataFrame([
            {"family_id": "F1", "representative_cost_item_name": "垂直运输费", "unit": "项", "本次召回样本数": 1, "本次召回工程包数": 1},
            {"family_id": "F2", "representative_cost_item_name": "垂直运输费", "unit": "项", "本次召回样本数": 2, "本次召回工程包数": 2},
        ])
        evidence = pd.DataFrame([
            {"family_id": "F1", "project_package_id": "P1"},
            {"family_id": "F2", "project_package_id": "P2"},
            {"family_id": "F2", "project_package_id": "P3"},
        ])
        no_match = types.SimpleNamespace(
            content={"decision": "no_explicit_match", "selected_option_id": ""}, usage={}, raw_content="{}"
        )
        with patch.object(query_estimate_llm, "request_llm_json_with_usage", return_value=no_match):
            selected, trace, _ = query_estimate_llm.select_final_options(
                "屋面漏水维修", plan, lookup, displays, families, [], evidence
            )
        self.assertEqual(selected.loc[0, "selected_option_id"], "D1-O02")
        self.assertEqual(trace.loc[0, "selection_reason"], "no_explicit_match_selected_by_sample_count")

        warnings = []
        with patch.object(query_estimate_llm, "request_llm_json_with_usage", side_effect=RuntimeError("down")):
            selected, trace, _ = query_estimate_llm.select_final_options(
                "屋面漏水维修", plan, lookup, displays, families, warnings, evidence
            )
        self.assertEqual(selected.loc[0, "selected_option_id"], "D1-O02")
        self.assertEqual(trace.loc[0, "selection_reason"], "llm_failed_selected_by_support")
        self.assertIn("option_selection_llm_failed_selected_by_support:0", warnings)

        invalid_results = [
            {"selected_option_id": "D1-O01"},
            {"decision": "no_explicit_match", "selected_option_id": "D1-O01"},
            {"decision": "explicit_match", "selected_option_id": "missing"},
            {"decision": "unknown", "selected_option_id": ""},
        ]
        for invalid in invalid_results:
            with self.subTest(invalid=invalid):
                response = types.SimpleNamespace(content=invalid, usage={}, raw_content="{}")
                warnings = []
                with patch.object(query_estimate_llm, "request_llm_json_with_usage", return_value=response):
                    selected, trace, _ = query_estimate_llm.select_final_options(
                        "屋面漏水维修", plan, lookup, displays, families, warnings, evidence
                    )
                self.assertEqual(selected.loc[0, "selected_option_id"], "D1-O02")
                self.assertEqual(trace.loc[0, "selection_reason"], "llm_failed_selected_by_support")
                self.assertIn("option_selection_llm_failed_selected_by_support:0", warnings)

    def test_representative_family_and_grouping_trace_record_selection_reasons(self):
        family_map = {
            "F1": pd.Series({"本次召回样本数": 1, "本次召回工程包数": 1}),
            "F2": pd.Series({"本次召回样本数": 3, "本次召回工程包数": 1}),
        }
        selected, reason = query_estimate_llm.choose_representative_family(
            {"family_ids": ["F1", "F2"]}, "F1", family_map
        )
        self.assertEqual((selected, reason), ("F2", "selected_option_family_sample_count_max"))

        trace = pd.DataFrame([{
            **{column: "" for column in query_estimate_llm.DISPLAY_OPTION_GROUPING_TRACE_COLUMNS},
            "display_id": "D1", "practice_option_id": "O2", "family_id": "F2",
            "option_sample_count": 3, "option_package_count": 1,
        }])
        selection = pd.DataFrame([{
            "display_id": "D1", "original_option_id": "O1", "selected_option_id": "O2",
            "original_family_id": "F1", "representative_family_id": "F2",
            "option_selection_decision": "no_explicit_match",
            "selection_reason": "no_explicit_match_selected_by_sample_count",
            "representative_selection_reason": reason,
        }])
        enriched = query_estimate_llm.apply_option_selection_to_grouping_trace(trace, selection).iloc[0]
        self.assertTrue(enriched["is_selected_option"])
        self.assertTrue(enriched["is_representative_family"])
        self.assertEqual(enriched["representative_selection_reason"], reason)

    def test_query_parse_args_explanations_default_false_and_flag_true(self):
        with patch.object(sys, "argv", ["query_cost_estimate_llm.py", "--text", "test"]):
            self.assertFalse(query_estimate_llm.parse_args().with_explanations)
        with patch.object(sys, "argv", ["query_cost_estimate_llm.py", "--text", "test", "--with-explanations"]):
            self.assertTrue(query_estimate_llm.parse_args().with_explanations)

    def test_optional_explanations_skip_or_call_existing_generator(self):
        scenario = query_estimate_llm.EstimateScenario("S001", 1, "", "", [])
        with patch.object(query_estimate_llm, "generate_final_explanation") as generator:
            output, success, error, prompt, trace = query_estimate_llm.generate_optional_final_explanation(
                False, "需求", scenario, pd.DataFrame(), []
            )
        generator.assert_not_called()
        self.assertEqual(output.scenario_summary, "当前结果根据保留清单形成一套历史常见维修组合，估价区间由组合内各清单的历史价格证据汇总形成，具体清单和价格见明细表。")
        self.assertTrue(success)
        self.assertEqual((error, prompt), ("", ""))
        self.assertEqual(trace["stage"], "final_explanation")

        expected = (scenario, True, "", "prompt", {"stage": "final_explanation"})
        with patch.object(query_estimate_llm, "generate_final_explanation", return_value=expected) as generator:
            actual = query_estimate_llm.generate_optional_final_explanation(
                True, "需求", scenario, pd.DataFrame()
            )
        generator.assert_called_once()
        self.assertEqual(actual, expected)

    def test_scenario_outputs_skip_unitless_item_without_blocking_normal_item(self):
        normal_item = query_estimate_llm.ScenarioItem(
            "package-1", "sid-waterproof", "ref-waterproof", "D1", "D1-O01",
            "D1-O01", "F1", "F1", "选用防水做法", {"type": "exact", "value": 10}, "按面积计算", 17,
            quantity_source="historical_median", quantity_explanation="没有用户数量",
            quantity_sample_count=3, quantity_minimum=5, quantity_median=10, quantity_maximum=15,
        )
        unitless_item = query_estimate_llm.ScenarioItem(
            "package-1", "sid-measures", "ref-measures", "D2", "D2-O01",
            "D2-O01", "F2", "F2", "其他措施费", {"type": "exact", "value": 1}, "按项计算",
        )
        scenario = query_estimate_llm.EstimateScenario(
            "S001", 1, "防水方案", "", [normal_item, unitless_item]
        )
        displays = pd.DataFrame([
            {"display_id": "D1", "unit": "m²", "practice_options": [
                {"practice_option_id": "D1-O01", "family_ids": ["F1"]},
            ]},
            {"display_id": "D2", "unit": "", "practice_options": [
                {"practice_option_id": "D2-O01", "family_ids": ["F2"]},
            ]},
        ])
        families = pd.DataFrame([
            {
                "family_id": "F1", "representative_cost_item_name": "屋面防水",
                "representative_project_description": "3mm SBS", "unit": "m²",
                "unit_normalized": "m²",
            },
            {
                "family_id": "F2", "representative_cost_item_name": "其他措施费",
                "representative_project_description": "", "unit": "", "unit_normalized": "",
            },
        ])
        price_stats = {
            "unit_price_p10": 100, "unit_price_median": 120, "unit_price_p90": 150,
            "labor_unit_price_p10": 10, "labor_unit_price_median": 12, "labor_unit_price_p90": 15,
            "machinery_unit_price_p10": 1, "machinery_unit_price_median": 2,
            "machinery_unit_price_p90": 3, "evidence_count": 2, "source_refs": "ref-1, ref-2",
            "expanded_evidence": pd.DataFrame([
                {"stable_sample_id": "price-1", "family_id": "F1", "normalized_signature": "sig"},
                {"stable_sample_id": "price-2", "family_id": "F1", "normalized_signature": "sig"},
            ]),
        }

        with patch.object(query_estimate_llm, "price_stats_for_option", return_value=price_stats) as price_lookup:
            output, price_evidence = query_estimate_llm.build_scenario_outputs(
                [scenario], displays, families, pd.DataFrame(), {17: ["F1", "F-extra"]}
            )

        price_lookup.assert_called_once()
        self.assertEqual(price_lookup.call_args.args[0]["practice_option_id"], "D1-O01")
        self.assertEqual(price_lookup.call_args.args[0]["family_ids"], ["F1", "F-extra"])
        self.assertEqual(len(output), 1)
        self.assertEqual(len(price_evidence), 2)
        self.assertEqual(output.loc[0, "价格证据样本数"], len(price_evidence))
        self.assertEqual(price_evidence["final_item_position"].tolist(), [17, 17])
        self.assertEqual(output.loc[0, "单位"], "m²")
        self.assertEqual(output.loc[0, "综合单价中位数"], 120)
        self.assertEqual(output.loc[0, "合价中位数"], 1200)
        self.assertEqual(output.loc[0, "工程量"], 10)
        self.assertNotIn("工程量来源", output.columns)
        self.assertEqual(output.loc[0, "工程量中位数"], 10)
        self.assertEqual(output.loc[0, "综合单价"], 120)
        self.assertEqual(output.loc[0, "暂估合价"], 1200)
        self.assertEqual(output.loc[0, "价格证据family"], "F1,F-extra")
        self.assertNotIn("工程量说明", output.columns)

    def test_estimate_scenario_columns_remove_quantity_count_and_follow_display_order(self):
        columns = query_estimate_llm.ESTIMATE_SCENARIO_COLUMNS
        self.assertNotIn("工程量样本数", columns)
        self.assertNotIn("工程量来源", columns)
        self.assertNotIn("工程量说明", columns)
        for required in ["价格证据样本数", "工程量最低值", "工程量中位数", "工程量最高值"]:
            self.assertIn(required, columns)
        self.assertEqual(columns[:5], ["清单名称", "项目特征", "单位", "工程量", "价格证据样本数"])
        self.assertEqual(columns[5:16], [
            "综合单价P10", "综合单价", "综合单价P90", "暂估合价",
            "其中包含人工费单价P10", "其中包含人工费单价中位数",
            "其中包含人工费单价P90", "其中包含机械费单价P10",
            "其中包含机械费单价中位数", "其中包含机械费单价P90",
            "工程量最低值",
        ])
        self.assertLess(columns.index("工程量最低值"), columns.index("工程量中位数"))
        self.assertLess(columns.index("工程量中位数"), columns.index("工程量最高值"))
        self.assertLess(columns.index("工程量最高值"), columns.index("来源样本"))
        self.assertEqual(columns[-6:], [
            "practice_option_id", "original_option_id", "original_family_id",
            "representative_family_id", "价格证据family", "display_id",
        ])

    def test_historical_quantity_display_uses_expanded_price_evidence_count(self):
        item = query_estimate_llm.ScenarioItem(
            "P1", "sid-selected", "selected-ref", "D1", "D1-O01", "D1-O01",
            "F1", "F1", "", {"type": "exact", "value": 720},
            "原说明包含4条同类历史样本", 5,
            quantity_source="historical_median", quantity_sample_count=4,
            quantity_minimum=600, quantity_median=720, quantity_maximum=900,
        )
        scenario = query_estimate_llm.EstimateScenario("S001", 1, "", "", [item])
        displays = pd.DataFrame([{
            "display_id": "D1", "unit": "m",
            "practice_options": [{"practice_option_id": "D1-O01", "family_ids": ["F1"]}],
        }])
        families = pd.DataFrame([self.price_family("F1", "sig-1", unit="m")])
        expanded_evidence = pd.DataFrame([
            {
                "stable_sample_id": f"price-{index}", "family_id": "F1",
                "normalized_signature": "sig-1", "source_ref": f"ref-{index}",
            }
            for index in range(61)
        ])
        price_stats = {
            "unit_price_p10": 10, "unit_price_median": 20, "unit_price_p90": 30,
            "labor_unit_price_p10": 1, "labor_unit_price_median": 2,
            "labor_unit_price_p90": 3, "machinery_unit_price_p10": 4,
            "machinery_unit_price_median": 5, "machinery_unit_price_p90": 6,
            "evidence_count": 61, "source_refs": "ref-0, ref-1",
            "expanded_evidence": expanded_evidence,
        }

        with patch.object(query_estimate_llm, "price_stats_for_option", return_value=price_stats):
            output, price_evidence = query_estimate_llm.build_scenario_outputs(
                [scenario], displays, families, pd.DataFrame()
            )

        self.assertEqual(output.loc[0, "工程量"], 720)
        self.assertEqual(output.loc[0, "工程量最低值"], 600)
        self.assertEqual(output.loc[0, "工程量中位数"], 720)
        self.assertEqual(output.loc[0, "工程量最高值"], 900)
        self.assertEqual(output.loc[0, "价格证据样本数"], 61)
        self.assertEqual(len(price_evidence), 61)
        self.assertNotIn("工程量说明", output.columns)
        self.assertEqual(output.loc[0, "综合单价"], 20)
        self.assertEqual(output.loc[0, "暂估合价"], 14400)
        self.assertEqual(output.loc[0, "display_id"], "D1")
        self.assertEqual(output.loc[0, "practice_option_id"], "D1-O01")
        self.assertEqual(output.loc[0, "original_option_id"], "D1-O01")
        self.assertEqual(output.loc[0, "representative_family_id"], "F1")

    def test_quantity_display_keeps_user_explicit_and_best_sample_fallback_reasons(self):
        displays = pd.DataFrame([{
            "display_id": "D1", "unit": "m",
            "practice_options": [{"practice_option_id": "D1-O01", "family_ids": ["F1"]}],
        }])
        families = pd.DataFrame([self.price_family("F1", "sig-1", unit="m")])
        price_stats = {
            "unit_price_p10": 10, "unit_price_median": 20, "unit_price_p90": 30,
            "labor_unit_price_p10": None, "labor_unit_price_median": None,
            "labor_unit_price_p90": None, "machinery_unit_price_p10": None,
            "machinery_unit_price_median": None, "machinery_unit_price_p90": None,
            "evidence_count": 1, "source_refs": "ref-1",
            "expanded_evidence": pd.DataFrame([{
                "stable_sample_id": "price-1", "family_id": "F1",
                "normalized_signature": "sig-1", "source_ref": "ref-1",
            }]),
        }
        cases = [
            ("user_explicit", False, "用户明确指定工程量为8m"),
            ("historical_median", True, "全库同类样本无有效中位数，暂采用最佳召回样本工程量。"),
        ]
        for source, fallback_used, reason in cases:
            item = query_estimate_llm.ScenarioItem(
                "P1", "sid-selected", "selected-ref", "D1", "D1-O01", "D1-O01",
                "F1", "F1", "", {"type": "exact", "value": 8}, reason, 1,
                quantity_source=source, quantity_fallback_used=fallback_used,
            )
            scenario = query_estimate_llm.EstimateScenario("S001", 1, "", "", [item])
            with self.subTest(source=source), patch.object(
                query_estimate_llm, "price_stats_for_option", return_value=price_stats
            ):
                output, _price_evidence = query_estimate_llm.build_scenario_outputs(
                    [scenario], displays, families, pd.DataFrame()
                )
            self.assertNotIn("工程量说明", output.columns)

    def price_family(self, family_id: str, signature: str, unit: str = "m²") -> dict[str, object]:
        return {
            "family_id": family_id,
            "normalized_signature": signature,
            "representative_cost_item_name": f"清单{family_id}",
            "representative_project_description": f"特征{family_id}",
            "unit": unit,
            "unit_normalized": unit,
        }

    def evidence_expansion_item(self) -> object:
        return query_estimate_llm.ScenarioItem(
            "P1", "selected", "selected-ref", "D1", "D1-O01", "D1-O01",
            "F1", "F1", "", {"type": "exact", "value": 1}, "", 3,
            quantity_source="historical_median",
        )

    def test_evidence_expansion_name_normalization_is_separate_from_display_normalization(self):
        plain = "更换曳引钢丝绳(含人工)"
        model = "更换曳引钢丝绳(P185012C000-01L221)(含人工)"

        self.assertEqual(query_estimate_llm.normalize_evidence_expansion_name(plain), "更换曳引钢丝绳")
        self.assertEqual(
            query_estimate_llm.normalize_evidence_expansion_name("更换曳引钢丝绳(含拆除及安装人工费)"),
            "更换曳引钢丝绳",
        )
        self.assertEqual(
            query_estimate_llm.normalize_evidence_expansion_name(model),
            "更换曳引钢丝绳(P185012C000-01L221)",
        )
        self.assertEqual(query_estimate_llm.normalize_display_name(plain), "更换曳引钢丝绳(含人工")

    def test_evidence_expansion_candidates_are_deduplicated_unit_exact_and_limited(self):
        item = self.evidence_expansion_item()
        option = {"practice_option_id": "D1-O01", "family_ids": ["F1"]}
        rows = [self.price_family("F1", "sig-1")]
        rows.append(self.price_family("F2", "sig-2", unit="㎡"))
        for index in range(3, 26):
            rows.append(self.price_family(f"F{index}", f"sig-{index}"))
        rows.insert(3, self.price_family("F3", "duplicate-signature"))

        original, target, candidates = query_estimate_llm.option_evidence_expansion_candidates(
            item, option, pd.DataFrame(rows)
        )

        self.assertEqual(original, ["F1"])
        self.assertEqual(target["unit"], "m²")
        self.assertEqual(len(candidates), 20)
        self.assertNotIn("F1", [row["family_id"] for row in candidates])
        self.assertNotIn("F2", [row["family_id"] for row in candidates])
        self.assertEqual([row["family_id"] for row in candidates[:2]], ["F3", "F4"])
        self.assertEqual(len({row["family_id"] for row in candidates}), 20)
        self.assertEqual(
            set(candidates[0]),
            {"family_id", "cost_item_name", "project_description", "unit"},
        )

    def evidence_tags(self, family_id: str, **overrides: str) -> dict[str, str]:
        tags = {
            "family_id": family_id,
            "object_action": "曳引钢丝绳更换",
            "thickness": "",
            "material": "",
            "level": "",
        }
        tags.update(overrides)
        return tags

    def test_evidence_expansion_validates_llm_output_strictly(self):
        parsed, accepted = query_estimate_llm.validate_option_evidence_expansion_result(
            {
                "families": [self.evidence_tags("F1"), self.evidence_tags("F2")],
                "accepted_family_ids": ["F2"],
            },
            "F1",
            ["F2"],
        )
        self.assertEqual(parsed["F1"]["object_action"], "曳引钢丝绳更换")
        self.assertEqual(accepted, ["F2"])

        invalid_results = [
            {},
            {"families": [], "accepted_family_ids": [], "extra": True},
            {"families": "not-an-array", "accepted_family_ids": []},
            {"families": [{"family_id": "F1"}], "accepted_family_ids": []},
            {"families": [{**self.evidence_tags("F1"), "extra": "x"}], "accepted_family_ids": []},
            {"families": [{**self.evidence_tags("F1"), "level": 1}], "accepted_family_ids": []},
            {"families": [self.evidence_tags("F1"), self.evidence_tags("F1")], "accepted_family_ids": []},
            {"families": [self.evidence_tags("F2")], "accepted_family_ids": []},
            {"families": [self.evidence_tags("F1")], "accepted_family_ids": []},
            {"families": [self.evidence_tags("F1"), self.evidence_tags("F9")], "accepted_family_ids": []},
            {"families": [self.evidence_tags("F1"), self.evidence_tags("F2")], "accepted_family_ids": "F2"},
            {"families": [self.evidence_tags("F1"), self.evidence_tags("F2")], "accepted_family_ids": ["F1"]},
            {"families": [self.evidence_tags("F1"), self.evidence_tags("F2")], "accepted_family_ids": ["F9"]},
            {"families": [self.evidence_tags("F1"), self.evidence_tags("F2")], "accepted_family_ids": ["F2", "F2"]},
        ]
        for result in invalid_results:
            with self.subTest(result=result), self.assertRaises(ValueError):
                query_estimate_llm.validate_option_evidence_expansion_result(
                    result, "F1", ["F2"]
                )

    def test_evidence_expansion_hard_conflicts_do_not_compare_object_action(self):
        tags = {
            "F1": {"object_action": "更换曳引钢丝绳", "thickness": "", "material": "", "level": ""},
            "F2": {"object_action": "曳引钢丝绳更换", "thickness": "", "material": "", "level": ""},
            "F3": {"object_action": "限速器钢丝绳更换", "thickness": "", "material": "", "level": ""},
        }
        self.assertEqual(
            query_estimate_llm.filter_option_evidence_hard_conflicts(
                ["F2", "F3"], "F1", tags, "m", {"F2": "m", "F3": "m"}
            ),
            ["F2", "F3"],
        )

    def test_evidence_expansion_hard_conflict_rules(self):
        tags = {
            "F1": {"object_action": "卷材防水新做", "thickness": "3mm", "material": "SBS", "level": "2层"},
            "F2": {"object_action": "卷材防水新做", "thickness": "4mm", "material": "SBS", "level": "2层"},
            "F3": {"object_action": "卷材防水新做", "thickness": "3mm", "material": "自粘卷材", "level": "2层"},
            "F4": {"object_action": "卷材防水新做", "thickness": "3mm", "material": "SBS", "level": "5层"},
            "F5": {"object_action": "卷材防水新做", "thickness": "3mm", "material": "SBS", "level": ""},
            "F6": {"object_action": "卷材防水新做", "thickness": "", "material": "", "level": "2层"},
            "F7": {"object_action": "卷材防水新做", "thickness": "3mm", "material": "SBS", "level": "2层"},
        }
        self.assertEqual(
            query_estimate_llm.filter_option_evidence_hard_conflicts(
                ["F2", "F3", "F4", "F5", "F6", "F7"],
                "F1",
                tags,
                "m²",
                {family_id: ("m" if family_id == "F7" else "m²") for family_id in tags if family_id != "F1"},
            ),
            ["F6"],
        )

    def test_evidence_expansion_llm_decision_controls_semantic_acceptance(self):
        fire_tags = {
            "F1": {"object_action": "消防广播主机更换", "thickness": "", "material": "", "level": ""},
            "F2": {"object_action": "更换消防广播主机", "thickness": "", "material": "", "level": ""},
            "F3": {"object_action": "消防报警主机更换", "thickness": "", "material": "", "level": ""},
            "F4": {"object_action": "多线盘更换", "thickness": "", "material": "", "level": ""},
        }
        self.assertEqual(
            query_estimate_llm.filter_option_evidence_hard_conflicts(
                ["F2"],
                "F1",
                fire_tags,
                "台",
                {"F2": "台", "F3": "台", "F4": "台"},
            ),
            ["F2"],
        )
        self.assertEqual(
            query_estimate_llm.filter_option_evidence_hard_conflicts(
                [], "F1", fire_tags, "台", {"F2": "台", "F3": "台", "F4": "台"}
            ),
            [],
        )

    def test_evidence_expansion_uses_selected_representative_without_reselection(self):
        original = self.evidence_expansion_item()
        item = query_estimate_llm.ScenarioItem(
            original.project_package_id, original.stable_sample_id, original.source_ref,
            original.display_id, original.practice_option_id, original.original_option_id,
            original.original_family_id, "F2", original.selection_reason, original.quantity,
            original.quantity_reason, original.item_position,
            quantity_source=original.quantity_source,
        )
        option = {"practice_option_id": "D1-O01", "family_ids": ["F1", "F2"]}
        families = pd.DataFrame([
            self.price_family("F1", "sig-1", unit="m"),
            self.price_family("F2", "sig-2", unit="台"),
            self.price_family("F3", "sig-3", unit="台"),
        ])

        original_family_ids, representative, candidates = (
            query_estimate_llm.option_evidence_expansion_candidates(item, option, families)
        )

        self.assertEqual(original_family_ids, ["F1", "F2"])
        self.assertEqual(representative["family_id"], "F2")
        self.assertEqual(representative["unit"], "台")
        self.assertEqual([candidate["family_id"] for candidate in candidates], ["F3"])

    def test_evidence_expansion_prompt_payload_is_minimal(self):
        representative = {
            "family_id": "F1", "cost_item_name": "清单F1",
            "project_description": "特征F1", "unit": "m²",
        }
        candidates = [{
            "family_id": "F2", "cost_item_name": "清单F2",
            "project_description": "特征F2", "unit": "m²",
        }]
        prompt = query_estimate_llm.build_option_evidence_expansion_prompt(representative, candidates)
        payload = json.loads(prompt.split("输入：\n", 1)[1])
        self.assertEqual(payload, {
            "representative_family": representative,
            "candidate_families": candidates,
        })
        for forbidden in ["price", "sample_count", "similarity", "source_refs", "accepted"]:
            self.assertNotIn(forbidden, json.dumps(payload, ensure_ascii=False))

    def test_evidence_expansion_success_updates_counts_and_deduplicates_samples(self):
        item = self.evidence_expansion_item()
        displays = pd.DataFrame([{
            "display_id": "D1", "unit": "m²",
            "practice_options": [{"practice_option_id": "D1-O01", "family_ids": ["F1"]}],
        }])
        families = pd.DataFrame([
            self.price_family("F1", "sig-1"),
            self.price_family("F2", "sig-2"),
            self.price_family("F3", "sig-3"),
            self.price_family("F4", "sig-4"),
        ])
        samples = pd.DataFrame([
            self.price_sample("sid-1", "sig-1", 10, source_ref="ref-1"),
            self.price_sample("sid-2", "sig-1", 20, source_ref="ref-2"),
            self.price_sample("sid-2", "sig-2", 999, source_ref="duplicate-ref"),
            self.price_sample("sid-3", "sig-2", 30, source_ref="ref-3"),
            self.price_sample("sid-4", "sig-3", 40, source_ref="ref-4"),
        ])
        response_content = {"families": [
            self.evidence_tags("F1", material="SBS改性沥青"),
            self.evidence_tags("F2", material="SBS改性沥青"),
            self.evidence_tags("F3", material="SBS改性沥青"),
            self.evidence_tags("F4", material="自粘卷材"),
        ], "accepted_family_ids": ["F4", "F3", "F2"]}
        response = types.SimpleNamespace(
            content=response_content, usage={}, raw_content=json.dumps(response_content, ensure_ascii=False)
        )

        with patch.object(query_estimate_llm, "request_llm_json_with_usage", return_value=response):
            expanded, sheet, traces = query_estimate_llm.expand_option_price_evidence_families(
                [item], displays, families, samples, []
            )

        self.assertEqual(expanded, {3: ["F1", "F2", "F3"]})
        self.assertEqual(sheet.loc[0, "新增family数"], 2)
        self.assertEqual(sheet.loc[0, "扩展后family数"], 3)
        self.assertEqual(sheet.loc[0, "原价格证据样本数"], 2)
        self.assertEqual(sheet.loc[0, "扩展后价格证据样本数"], 4)
        self.assertEqual(sheet.loc[0, "新增family_ids"], "F2,F3")
        self.assertEqual(traces[0]["stage"], "option_evidence_expansion")
        self.assertIn("llm_accepted=3; accepted_after_guard=2", traces[0]["input_summary"])
        self.assertFalse(traces[0]["fallback"])

    def test_evidence_expansion_threshold_uses_original_evidence_count(self):
        item = self.evidence_expansion_item()
        displays = pd.DataFrame([{
            "display_id": "D1", "unit": "m²",
            "practice_options": [{"practice_option_id": "D1-O01", "family_ids": ["F1"]}],
        }])
        families = pd.DataFrame([
            self.price_family("F1", "sig-1"), self.price_family("F2", "sig-2"),
        ])
        samples = pd.DataFrame([
            self.price_sample(f"sid-{index}", "sig-1", index)
            for index in range(1, 11)
        ])

        with patch.object(query_estimate_llm, "request_llm_json_with_usage") as llm:
            expanded, sheet, traces = query_estimate_llm.expand_option_price_evidence_families(
                [item], displays, families, samples, []
            )

        llm.assert_not_called()
        self.assertEqual(expanded, {3: ["F1"]})
        self.assertEqual(sheet.loc[0, "候选family数"], 0)
        self.assertEqual(sheet.loc[0, "原价格证据样本数"], 10)
        self.assertEqual(sheet.loc[0, "扩展后价格证据样本数"], 10)
        self.assertEqual(traces[0]["parsed_status"], "skipped")
        self.assertIn("original_evidence=10; skipped=threshold", traces[0]["input_summary"])
        self.assertFalse(traces[0]["fallback"])

    def test_evidence_expansion_nine_samples_calls_llm_once(self):
        item = self.evidence_expansion_item()
        displays = pd.DataFrame([{
            "display_id": "D1", "unit": "m²",
            "practice_options": [{"practice_option_id": "D1-O01", "family_ids": ["F1"]}],
        }])
        families = pd.DataFrame([
            self.price_family("F1", "sig-1"), self.price_family("F2", "sig-2"),
        ])
        samples = pd.DataFrame([
            self.price_sample(f"sid-{index}", "sig-1", index)
            for index in range(1, 10)
        ] + [self.price_sample("sid-10", "sig-2", 10)])
        response_content = {
            "families": [
                self.evidence_tags("F1", object_action="更换曳引钢丝绳"),
                self.evidence_tags("F2", object_action="曳引钢丝绳更换"),
            ],
            "accepted_family_ids": ["F2"],
        }
        response = types.SimpleNamespace(content=response_content, usage={}, raw_content="")

        with patch.object(
            query_estimate_llm, "request_llm_json_with_usage", return_value=response
        ) as llm:
            expanded, _sheet, _traces = query_estimate_llm.expand_option_price_evidence_families(
                [item], displays, families, samples, []
            )

        llm.assert_called_once()
        self.assertEqual(expanded, {3: ["F1", "F2"]})

    def test_evidence_expansion_no_candidates_is_normal_skip(self):
        item = self.evidence_expansion_item()
        displays = pd.DataFrame([{
            "display_id": "D1", "unit": "m²",
            "practice_options": [{"practice_option_id": "D1-O01", "family_ids": ["F1"]}],
        }])
        families = pd.DataFrame([self.price_family("F1", "sig-1")])
        samples = pd.DataFrame([self.price_sample("sid-1", "sig-1", 10)])
        warnings = []

        with patch.object(query_estimate_llm, "request_llm_json_with_usage") as llm:
            expanded, sheet, traces = query_estimate_llm.expand_option_price_evidence_families(
                [item], displays, families, samples, warnings
            )

        llm.assert_not_called()
        self.assertEqual(expanded, {3: ["F1"]})
        self.assertEqual(sheet.loc[0, "候选family数"], 0)
        self.assertEqual(traces[0]["parsed_status"], "skipped")
        self.assertIn("skipped=no_candidates", traces[0]["input_summary"])
        self.assertFalse(traces[0]["fallback"])
        self.assertEqual(warnings, [])

    def test_evidence_expansion_invalid_representative_falls_back_without_reselection(self):
        original = self.evidence_expansion_item()
        item = query_estimate_llm.ScenarioItem(
            original.project_package_id, original.stable_sample_id, original.source_ref,
            original.display_id, original.practice_option_id, original.original_option_id,
            original.original_family_id, "F2", original.selection_reason, original.quantity,
            original.quantity_reason, original.item_position,
            quantity_source=original.quantity_source,
        )
        displays = pd.DataFrame([{
            "display_id": "D1", "display_name": "清单F1", "unit": "m²",
            "practice_options": [{"practice_option_id": "D1-O01", "family_ids": ["F1"]}],
        }])
        families = pd.DataFrame([
            self.price_family("F1", "sig-1"), self.price_family("F2", "sig-2"),
        ])
        samples = pd.DataFrame([self.price_sample("sid-1", "sig-1", 10)])
        warnings = []

        with patch.object(query_estimate_llm, "request_llm_json_with_usage") as llm:
            expanded, _sheet, traces = query_estimate_llm.expand_option_price_evidence_families(
                [item], displays, families, samples, warnings
            )

        llm.assert_not_called()
        self.assertEqual(expanded, {3: ["F1"]})
        self.assertEqual(warnings, ["option_evidence_expansion_failed:item_position=3"])
        self.assertTrue(traces[0]["fallback"])

    def test_evidence_expansion_failure_falls_back_without_blocking(self):
        item = self.evidence_expansion_item()
        displays = pd.DataFrame([{
            "display_id": "D1", "unit": "m²",
            "practice_options": [{"practice_option_id": "D1-O01", "family_ids": ["F1"]}],
        }])
        families = pd.DataFrame([
            self.price_family("F1", "sig-1"), self.price_family("F2", "sig-2"),
        ])
        samples = pd.DataFrame([self.price_sample("sid-1", "sig-1", 10)])
        warnings = []

        with patch.object(query_estimate_llm, "request_llm_json_with_usage", side_effect=RuntimeError("down")):
            expanded, sheet, traces = query_estimate_llm.expand_option_price_evidence_families(
                [item], displays, families, samples, warnings
            )

        self.assertEqual(expanded, {3: ["F1"]})
        self.assertEqual(sheet.loc[0, "新增family数"], 0)
        self.assertEqual(sheet.loc[0, "扩展后价格证据样本数"], 1)
        self.assertEqual(warnings, ["option_evidence_expansion_failed:item_position=3"])
        self.assertTrue(traces[0]["fallback"])

    def price_sample(
        self,
        stable_sample_id: str,
        signature: str,
        unit_price: object,
        *,
        unit: str = "m²",
        source_ref: str = "",
    ) -> dict[str, object]:
        index = int(re.sub(r"\D", "", stable_sample_id) or 1)
        return {
            "stable_sample_id": stable_sample_id,
            "normalized_signature": signature,
            "unit": unit,
            "unit_normalized": unit,
            "unit_price": unit_price,
            "labor_unit_price": index,
            "machinery_unit_price": index * 2,
            "source_ref": source_ref,
            "project_key": f"project-{index}",
            "item_row_id": f"row-{index}",
        }

    def test_option_price_stats_expand_single_family_from_all_samples(self):
        option = {"practice_option_id": "D001-O01", "family_ids": ["F001"]}
        families = pd.DataFrame([self.price_family("F001", "sig-1")])
        samples = pd.DataFrame([
            self.price_sample(f"sid-{index}", "sig-1", price, source_ref=f"ref-{index}")
            for index, price in enumerate([10, 20, 30, 40, 50], start=1)
        ])

        stats = query_estimate_llm.price_stats_for_option(
            option, pd.Series({"unit": "m²"}), families, samples
        )

        self.assertEqual(stats["evidence_count"], 5)
        self.assertEqual(
            (stats["unit_price_p10"], stats["unit_price_median"], stats["unit_price_p90"]),
            (10.0, 30.0, 50.0),
        )
        self.assertEqual(
            (stats["labor_unit_price_p10"], stats["labor_unit_price_median"], stats["labor_unit_price_p90"]),
            (1.0, 3.0, 5.0),
        )

    def test_option_price_stats_uses_nearest_sample_quantiles_not_min_max(self):
        option = {"practice_option_id": "D001-O01", "family_ids": ["F001"]}
        families = pd.DataFrame([self.price_family("F001", "sig-1")])
        samples = pd.DataFrame([
            self.price_sample(f"sid-{index}", "sig-1", index)
            for index in range(1, 12)
        ])

        stats = query_estimate_llm.price_stats_for_option(
            option, pd.Series({"unit": "m²"}), families, samples
        )

        self.assertEqual(
            (stats["unit_price_p10"], stats["unit_price_median"], stats["unit_price_p90"]),
            (2.0, 6.0, 10.0),
        )
        self.assertIn(stats["unit_price_p10"], samples["unit_price"].tolist())
        self.assertIn(stats["unit_price_p90"], samples["unit_price"].tolist())

    def test_option_price_stats_expand_only_within_location_and_time_constraints(self):
        option = {"practice_option_id": "D001-O01", "family_ids": ["F001"]}
        families = pd.DataFrame([self.price_family("F001", "sig-1")])
        samples = pd.DataFrame([
            {
                **self.price_sample("sid-1", "sig-1", 100, source_ref="inside-1"),
                "location": "浙江省嘉兴市", "consultation_time": "2025-06-01",
            },
            {
                **self.price_sample("sid-2", "sig-1", 200, source_ref="inside-2"),
                "location": "浙江省嘉兴市", "consultation_time": "2025-12-31",
            },
            {
                **self.price_sample("sid-3", "sig-1", 1000, source_ref="outside-location"),
                "location": "上海市", "consultation_time": "2025-06-01",
            },
            {
                **self.price_sample("sid-4", "sig-1", 2000, source_ref="outside-time"),
                "location": "浙江省嘉兴市", "consultation_time": "2024-12-31",
            },
        ])
        embeddings = np.zeros((len(samples), 1), dtype=np.float32)

        constrained_mask = query_estimate_llm.build_constraint_mask(
            samples, "浙江省嘉兴市", "2025-01-01", "2025-12-31"
        )
        candidate_samples, _candidate_embeddings = query_estimate_llm.filter_rows_and_embeddings(
            samples, embeddings, constrained_mask, "清单样本"
        )
        constrained_stats = query_estimate_llm.price_stats_for_option(
            option, pd.Series({"unit": "m²"}), families, candidate_samples
        )

        self.assertEqual(constrained_stats["evidence_count"], 2)
        self.assertEqual(constrained_stats["source_refs"], "inside-1, inside-2")
        self.assertEqual(
            (constrained_stats["unit_price_p10"], constrained_stats["unit_price_median"], constrained_stats["unit_price_p90"]),
            (100.0, 150.0, 200.0),
        )

        unconstrained_mask = query_estimate_llm.build_constraint_mask(samples, "", "", "")
        all_candidate_samples, _all_candidate_embeddings = query_estimate_llm.filter_rows_and_embeddings(
            samples, embeddings, unconstrained_mask, "清单样本"
        )
        unconstrained_stats = query_estimate_llm.price_stats_for_option(
            option, pd.Series({"unit": "m²"}), families, all_candidate_samples
        )

        self.assertEqual(unconstrained_stats["evidence_count"], 4)
        self.assertEqual(unconstrained_stats["unit_price_p90"], 2000.0)
        self.assertIn("outside-location", unconstrained_stats["source_refs"])
        self.assertIn("outside-time", unconstrained_stats["source_refs"])

    def test_option_price_stats_expand_multiple_families(self):
        option = {"practice_option_id": "D001-O01", "family_ids": ["F001", "F002"]}
        families = pd.DataFrame([
            self.price_family("F001", "sig-1"),
            self.price_family("F002", "sig-2"),
        ])
        samples = pd.DataFrame([
            self.price_sample(f"sid-{index}", "sig-1" if index <= 5 else "sig-2", index, source_ref=f"ref-{index}")
            for index in range(1, 9)
        ])

        stats = query_estimate_llm.price_stats_for_option(
            option, pd.Series({"unit": "m²"}), families, samples
        )

        self.assertEqual(stats["evidence_count"], 8)
        self.assertEqual(stats["unit_price_median"], 4.5)

    def test_option_price_stats_limit_sources_without_limiting_count_or_statistics(self):
        option = {"practice_option_id": "D001-O01", "family_ids": ["F001"]}
        families = pd.DataFrame([self.price_family("F001", "sig-1")])
        samples = pd.DataFrame([
            self.price_sample(f"sid-{index}", "sig-1", index, source_ref=f"ref-{index}")
            for index in range(1, 16)
        ])

        stats = query_estimate_llm.price_stats_for_option(
            option, pd.Series({"unit": "m²"}), families, samples
        )

        self.assertEqual(stats["evidence_count"], 15)
        self.assertEqual(stats["unit_price_median"], 8.0)
        self.assertEqual(stats["machinery_unit_price_p90"], 28.0)
        self.assertEqual(len(stats["source_refs"].split(", ")), 10)

    def test_option_price_stats_deduplicate_shared_signature_and_preserve_or_generate_sources(self):
        option = {"practice_option_id": "D001-O01", "family_ids": ["F001", "F002"]}
        families = pd.DataFrame([
            self.price_family("F001", "shared"),
            self.price_family("F002", "shared"),
        ])
        samples = pd.DataFrame([
            self.price_sample("sid-1", "shared", 10, source_ref="existing-ref"),
            self.price_sample("sid-2", "shared", 20),
        ])

        stats = query_estimate_llm.price_stats_for_option(
            option, pd.Series({"unit": "m²"}), families, samples
        )

        self.assertEqual(stats["evidence_count"], 2)
        self.assertEqual(stats["source_refs"], "existing-ref, project-2::row-2")

    def test_option_price_stats_reject_invalid_family_mappings_and_missing_signatures(self):
        display = pd.Series({"unit": "m²"})
        samples = pd.DataFrame([self.price_sample("sid-1", "sig-1", 10)])
        cases = [
            (pd.DataFrame([self.price_family("F002", "sig-1")]), "不存在"),
            (pd.DataFrame([self.price_family("F001", "sig-1"), self.price_family("F001", "sig-1")]), "不唯一"),
            (pd.DataFrame([self.price_family("F001", "")]), "为空"),
            (pd.DataFrame([self.price_family("F001", "missing")]), "无匹配"),
        ]
        option = {"practice_option_id": "D001-O01", "family_ids": ["F001"]}
        for families, expected in cases:
            with self.subTest(expected=expected), self.assertRaisesRegex(ValueError, f"D001-O01.*F001|F001.*{expected}") as raised:
                query_estimate_llm.price_stats_for_option(option, display, families, samples)
            self.assertIn("F001", str(raised.exception))
            self.assertIn("D001-O01", str(raised.exception))

    def test_option_price_stats_falls_back_to_source_ref_and_rejects_incompatible_units(self):
        option = {"practice_option_id": "D001-O01", "family_ids": ["F001"]}
        families = pd.DataFrame([self.price_family("F001", "sig-1")])
        source_deduplicated = query_estimate_llm.price_stats_for_option(
            option,
            pd.Series({"unit": "m²"}),
            families,
            pd.DataFrame([
                self.price_sample("", "sig-1", 10, source_ref="same-ref"),
                self.price_sample("", "sig-1", 20, source_ref="same-ref"),
            ]),
        )
        self.assertEqual(source_deduplicated["evidence_count"], 1)
        with self.assertRaisesRegex(ValueError, "单位不兼容"):
            query_estimate_llm.price_stats_for_option(
                option,
                pd.Series({"unit": "m²"}),
                families,
                pd.DataFrame([self.price_sample("sid-1", "sig-1", 10, unit="m")]),
            )

        duplicate_stats = query_estimate_llm.price_stats_for_option(
            option,
            pd.Series({"unit": "m²"}),
            families,
            pd.DataFrame([
                self.price_sample("sid-1", "sig-1", 10),
                self.price_sample("sid-1", "sig-1", 20),
            ]),
        )
        self.assertEqual(duplicate_stats["evidence_count"], 1)
        self.assertEqual(duplicate_stats["unit_price_median"], 10)

    def test_option_price_stats_do_not_fallback_when_comprehensive_prices_are_empty(self):
        option = {"practice_option_id": "D001-O01", "family_ids": ["F001"]}
        families = pd.DataFrame([{
            **self.price_family("F001", "sig-1"),
            "本次召回综合单价中位数": 999,
        }])
        samples = pd.DataFrame([self.price_sample("sid-1", "sig-1", "")])

        stats = query_estimate_llm.price_stats_for_option(
            option, pd.Series({"unit": "m²"}), families, samples
        )
        with self.assertRaisesRegex(ValueError, "unit_price"):
            query_estimate_llm.validate_price_stats(stats, "selected-id", "D001-O01")

    def test_quantity_validation(self):
        self.assertEqual(query_estimate_llm.validate_quantity({"type": "exact", "value": 10}), {"type": "exact", "value": 10.0})
        for quantity in [
            {"type": "exact", "value": 0},
            {"type": "exact", "value": -1},
            {"type": "range", "min": 20, "max": 10},
            {"type": "unknown"},
        ]:
            with self.assertRaises(ValueError):
                query_estimate_llm.validate_quantity(quantity)

    def test_parse_info_records_exact_only_quantity_and_explanation_config(self):
        self.assertEqual(query_estimate_llm.quantity_determination_status({"error_message": ""}), "success")
        self.assertEqual(query_estimate_llm.quantity_determination_status({"error_message": "invalid"}), "failed")
        self.assertEqual(
            query_estimate_llm.quantity_determination_status({"error_message": "invalid", "fallback": True}),
            "fallback",
        )
        trace = {"prompt_chars": 10, "prompt_tokens": 4, "completion_tokens": 2}
        parse_info = query_estimate_llm.build_parse_info(
            rewrite=query_estimate_llm.QueryRewrite("屋面", "屋面工程", "屋面防水", "浙江省嘉兴市", "2025-01-01", "2025-12-31", [], True),
            top_packages=20,
            top_items=300,
            max_packages_per_cache_subject=1,
            package_weight_temperature=0.1,
            evidence_package_universe_count=20,
            package_evidence_weight_count=20,
            package_evidence_weight_sum=1.0,
            meta={},
            sample_count=100,
            package_count=10,
            project_packages_before_constraint=10,
            project_packages_after_constraint=4,
            samples_before_constraint=100,
            samples_after_constraint=40,
            invalid_sample_consultation_time_count=0,
            invalid_project_package_consultation_time_count=0,
            retrieved_evidence_item_row_count=50,
            evidence_item_row_count=50,
            candidate_family_count=8,
            candidate_display_group_count=4,
            matched_project_example_count=3,
            matched_project_example_item_count=30,
            display_option_grouping_display_count=4,
            display_option_grouping_trace=trace,
            display_option_grouping_fallback=False,
            display_option_grouping_error="",
            display_option_grouping_meta={},
            selected_project_package_id="PKG-1",
            selected_item_count=2,
            selected_exact_quantity_count=1,
            selected_range_quantity_count=0,
            range_selection_trace=trace,
            range_selection_error="fallback reason",
            quantity_trace=trace,
            final_explanation_trace=trace,
            final_explanation_error="invalid explanation",
            output_path=Path("query.xlsx"),
            started_at=query_estimate_llm.datetime.now(),
            index_dir=Path("embeddings"),
            include_debug_text=True,
            display_option_grouping_prompt="grouping",
            range_selection_prompt="range",
            quantity_prompt="quantity",
            final_explanation_prompt="explanation",
            with_explanations=False,
            warnings=["final_explanation_failed"],
        )
        values = dict(parse_info.values.tolist())

        self.assertEqual(values["selected_project_package_id"], "PKG-1")
        self.assertEqual(values["query_location"], "浙江省嘉兴市")
        self.assertEqual(values["project_packages_after_constraint"], 4)
        self.assertEqual(values["samples_after_constraint"], 40)
        self.assertEqual(values["range_selection_status"], "fallback_full_project")
        self.assertEqual(values["quantity_determination_status"], "success")
        self.assertEqual(values["with_explanations"], False)
        self.assertEqual(values["final_explanation_status"], "skipped")
        self.assertIn("invalid explanation", values["final_explanation LLM error"])
        self.assertEqual(values["range_selection_prompt_preview"], "range")
        self.assertEqual(values["quantity_determination_prompt_preview"], "quantity")

    def test_estimate_summary_uses_final_columns_and_sums_all_items(self):
        scenario = query_estimate_llm.EstimateScenario(
            "S001", 1, "维修方案", "方案说明", [], "确认现场"
        )
        estimate_scenarios = pd.DataFrame([
            {"清单名称": "清单A", "工程量": 2, "综合单价P10": 5, "综合单价中位数": 10, "综合单价P90": 15},
            {"清单名称": "清单B", "工程量": 0.5, "综合单价P10": 3, "综合单价中位数": 5, "综合单价P90": 7},
        ])
        price_evidence = pd.DataFrame([
            {"project_key": "P1", "project_package_id": "PKG1", "工程名称": "重复名称"},
            {"project_key": "P1", "project_package_id": "PKG1", "工程名称": "重复名称"},
            {"project_key": "P2", "project_package_id": "PKG2", "工程名称": "重复名称"},
            {"project_key": "", "project_package_id": "PKG3", "工程名称": "独立名称"},
        ])

        summary = query_estimate_llm.build_estimate_summary("原始问题", [scenario], estimate_scenarios, price_evidence)

        self.assertEqual(summary.columns.tolist(), query_estimate_llm.ESTIMATE_SUMMARY_COLUMNS)
        self.assertEqual(
            summary.loc[0, ["合价P10", "合价中位数", "合价P90"]].tolist(),
            [11.5, 22.5, 33.5],
        )
        self.assertEqual(summary.loc[0, "用户问题"], "原始问题")
        self.assertEqual(summary.loc[0, "参考项目数"], 2)
        self.assertEqual(summary.loc[0, "参考样本数"], 4)
        for removed in ["方案顺序", "方案编号", "是否推荐方案", "与其他方案的核心差异", "合价最低值", "合价最高值"]:
            self.assertNotIn(removed, summary.columns)

    def test_customer_display_filter_keeps_threshold_and_complete_debug_evidence(self):
        items = [
            query_estimate_llm.ScenarioItem(
                "PKG", f"sid-{position}", f"ref-{position}", f"D{position}", f"O{position}",
                f"O{position}", f"F{position}", f"F{position}", "", {"type": "exact", "value": 1}, "", position,
            )
            for position in [1, 2, 3]
        ]
        scenario = query_estimate_llm.EstimateScenario("S001", 1, "", "", items)
        rows = pd.DataFrame([
            {"final_item_position": 1, "清单名称": "A", "价格证据样本数": 2},
            {"final_item_position": 2, "清单名称": "B", "价格证据样本数": 3},
            {"final_item_position": 3, "清单名称": "C", "价格证据样本数": 10},
        ])
        all_evidence = pd.DataFrame([
            {"final_item_position": position, "stable_sample_id": f"e-{position}-{index}"}
            for position, count in [(1, 2), (2, 3), (3, 10)]
            for index in range(count)
        ])

        filtered_scenario, display_rows, display_evidence = query_estimate_llm.filter_customer_display_outputs(
            scenario, rows, all_evidence
        )

        self.assertEqual(display_rows["清单名称"].tolist(), ["B", "C"])
        self.assertEqual([item.item_position for item in filtered_scenario.items], [2, 3])
        self.assertEqual(len(display_evidence), 13)
        self.assertEqual(len(all_evidence), 15)
        self.assertNotIn("final_item_position", display_rows.columns)

    def test_summary_component_amounts_missing_values_and_reference_project_fallback(self):
        rows = pd.DataFrame([
            {
                "清单名称": "A", "工程量": 10,
                "综合单价P10": 1, "综合单价中位数": 2, "综合单价P90": 3,
                "其中包含人工费单价P10": 4, "其中包含人工费单价中位数": 5, "其中包含人工费单价P90": 6,
                "其中包含机械费单价P10": "", "其中包含机械费单价中位数": "", "其中包含机械费单价P90": "",
            },
            {
                "清单名称": "A", "工程量": 2,
                "综合单价P10": 10, "综合单价中位数": 20, "综合单价P90": 30,
                "其中包含人工费单价P10": "", "其中包含人工费单价中位数": 20, "其中包含人工费单价P90": "",
                "其中包含机械费单价P10": "", "其中包含机械费单价中位数": "", "其中包含机械费单价P90": "",
            },
        ])
        evidence = pd.DataFrame([
            {"project_key": "", "project_package_id": "PKG1"},
            {"project_key": "", "project_package_id": "PKG1"},
            {"project_key": "", "project_package_id": "PKG2"},
        ])
        scenario = query_estimate_llm.EstimateScenario("S001", 1, "方案", "说明", [], "确认")

        summary = query_estimate_llm.build_estimate_summary("问题", [scenario], rows, evidence).iloc[0]

        self.assertEqual(summary["主要施工内容"], "A")
        self.assertEqual(summary["计价项目数"], 2)
        self.assertEqual(summary["参考项目数"], 2)
        self.assertEqual(summary["其中包含人工费中位数"], 90)
        self.assertEqual(summary["其中包含人工费P10"], 40)
        self.assertEqual(summary["其中包含人工费P90"], 60)
        self.assertEqual(summary["其中包含机械费中位数"], "")

    def test_final_explanation_prompt_and_strict_result_validation(self):
        final_items = [{
            "cost_item_name": "B", "project_description": "3mm", "unit": "m²",
            "quantity": 10, "unit_price": 20, "estimated_amount": 200,
        }]
        prompt = query_estimate_llm.build_final_explanation_prompt("原始问题", final_items, 100, 200, 300)
        payload = json.loads(prompt.split("输入：\n", 1)[1])
        self.assertEqual(set(payload), {"user_query", "final_items", "total_price"})
        self.assertEqual(payload["final_items"], final_items)
        self.assertNotIn("evidence_count", prompt.split("输入：\n", 1)[1])

        scenario = query_estimate_llm.EstimateScenario("S001", 1, "", "", [])
        valid = {"scenario_name": "方案", "scenario_summary": "说明", "site_confirmation": "确认"}
        parsed = query_estimate_llm.parse_final_explanation_result(valid, scenario)
        self.assertEqual(parsed.site_confirmation, "确认")
        invalid_results = [
            {**valid, "extra": "x"},
            {"scenario_name": "方案", "scenario_summary": "说明"},
            {**valid, "scenario_name": ""},
            {**valid, "scenario_summary": []},
            {**valid, "scenario_summary": "**Markdown**"},
            [valid],
        ]
        for invalid in invalid_results:
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                query_estimate_llm.parse_final_explanation_result(invalid, scenario)

    def test_invalid_final_explanation_uses_filtered_fallback_without_failing_query(self):
        item = query_estimate_llm.ScenarioItem(
            "PKG", "sid-b", "ref-b", "D2", "O2", "O2", "F2", "F2", "",
            {"type": "exact", "value": 2}, "", 2,
        )
        scenario = query_estimate_llm.EstimateScenario("S001", 1, "", "", [item])
        rows = pd.DataFrame([{
            "清单名称": "清单B", "项目特征": "特征B", "单位": "m", "工程量": 2,
            "综合单价P10": 10, "综合单价中位数": 20, "综合单价P90": 30,
        }])
        response = types.SimpleNamespace(
            content={"scenario_name": "方案", "scenario_summary": "说明", "site_confirmation": "确认", "extra": "x"},
            usage={}, raw_content="{}",
        )
        warnings = []

        with patch.object(query_estimate_llm, "request_llm_json_with_usage", return_value=response):
            output, success, error, prompt, trace = query_estimate_llm.generate_final_explanation(
                "原始问题", scenario, rows, warnings
            )

        self.assertFalse(success)
        self.assertIn("顶层字段非法", error)
        self.assertEqual(output.scenario_name, "清单B维修组合参考方案")
        self.assertNotIn("清单A", prompt)
        self.assertIn("final_explanation_failed", warnings)
        self.assertEqual(trace["parsed_status"], "failed")

    def test_all_filtered_summary_is_one_fixed_row_and_no_amounts(self):
        base = query_estimate_llm.EstimateScenario("S001", 1, "", "", [])
        scenario = query_estimate_llm.insufficient_evidence_scenario(base)
        rows = pd.DataFrame(columns=query_estimate_llm.ESTIMATE_SCENARIO_COLUMNS)
        evidence = pd.DataFrame(columns=query_estimate_llm.PRICE_EVIDENCE_ITEM_COLUMNS)

        summary = query_estimate_llm.build_estimate_summary("原始问题", [scenario], rows, evidence)

        self.assertEqual(len(summary), 1)
        self.assertEqual(summary.loc[0, "方案名称"], "历史价格证据不足")
        self.assertEqual(summary.loc[0, "计价项目数"], 0)
        self.assertEqual(summary.loc[0, "参考项目数"], 0)
        self.assertEqual(summary.loc[0, "参考样本数"], 0)
        self.assertEqual(summary.loc[0, "合价中位数"], "")

        with patch.object(query_estimate_llm, "generate_optional_final_explanation") as generator:
            explained, success, error, prompt, trace = query_estimate_llm.generate_customer_explanation(
                True, "原始问题", base, rows
            )
        generator.assert_not_called()
        self.assertTrue(success)
        self.assertEqual((error, prompt), ("", ""))
        self.assertEqual(explained.scenario_name, "历史价格证据不足")
        self.assertEqual(trace["parsed_status"], "skipped")

    def test_workbook_includes_new_columns_and_two_stage_trace(self):
        result = query_estimate_llm.QueryResult(
            rewrite=query_estimate_llm.QueryRewrite("屋面", "屋面工程", "屋面防水", "", "", "", [], True),
            estimate_summary=pd.DataFrame(columns=query_estimate_llm.ESTIMATE_SUMMARY_COLUMNS),
            estimate_scenarios=pd.DataFrame(columns=query_estimate_llm.ESTIMATE_SCENARIO_COLUMNS),
            matched_project_packages=pd.DataFrame(columns=query_estimate_llm.MATCHED_PROJECT_PACKAGE_COLUMNS),
            candidate_families=pd.DataFrame(columns=query_estimate_llm.CANDIDATE_FAMILY_COLUMNS),
            candidate_display_groups=pd.DataFrame(columns=query_estimate_llm.CANDIDATE_DISPLAY_GROUP_COLUMNS),
            package_evidence_weights=pd.DataFrame(columns=query_estimate_llm.PACKAGE_EVIDENCE_WEIGHT_COLUMNS),
            display_group_families=pd.DataFrame(columns=query_estimate_llm.DISPLAY_GROUP_FAMILY_COLUMNS),
            display_option_grouping_trace=pd.DataFrame(columns=query_estimate_llm.DISPLAY_OPTION_GROUPING_TRACE_COLUMNS),
            option_selection_trace=pd.DataFrame(columns=query_estimate_llm.OPTION_SELECTION_TRACE_COLUMNS),
            matched_project_examples=pd.DataFrame(columns=query_estimate_llm.MATCHED_PROJECT_EXAMPLE_COLUMNS),
            evidence_items=pd.DataFrame(columns=query_estimate_llm.EVIDENCE_ITEM_COLUMNS),
            option_evidence_expansion=pd.DataFrame(
                columns=query_estimate_llm.OPTION_EVIDENCE_EXPANSION_COLUMNS
            ),
            price_evidence_items=pd.DataFrame(columns=query_estimate_llm.PRICE_EVIDENCE_ITEM_COLUMNS),
            parse_info=pd.DataFrame([{"字段": "final_explanation_status", "值": "failed"}]),
            llm_trace=pd.DataFrame(
                [{"stage": stage} for stage in ["query_rewrite_for_embedding", "display_option_grouping", "range_selection", "quantity_determination", "final_explanation"]],
                columns=query_estimate_llm.LLM_TRACE_COLUMNS,
            ),
            success=False,
            error_message="explanation failed",
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "result.xlsx"
            query_estimate_llm.write_query_result_workbook(path, result)
            workbook = openpyxl.load_workbook(path, data_only=True)
            scenario_headers = [cell.value for cell in workbook["estimate_scenarios"][1]]
            expansion_headers = [cell.value for cell in workbook["option_evidence_expansion"][1]]
            price_evidence_headers = [cell.value for cell in workbook["price_evidence_items"][1]]
            trace_stages = [workbook["llm_trace"].cell(row=row, column=1).value for row in range(2, 7)]
            sheetnames = workbook.sheetnames
            self.assertNotIn("range_selection", workbook.sheetnames)
            workbook.close()

        self.assertEqual(
            scenario_headers,
            query_estimate_llm.ESTIMATE_SCENARIO_COLUMNS,
        )
        self.assertEqual(price_evidence_headers, query_estimate_llm.PRICE_EVIDENCE_ITEM_COLUMNS)
        self.assertEqual(expansion_headers, query_estimate_llm.OPTION_EVIDENCE_EXPANSION_COLUMNS)
        self.assertEqual(
            sheetnames[1:4],
            ["estimate_scenarios", "option_evidence_expansion", "price_evidence_items"],
        )
        self.assertEqual(scenario_headers[-1], "display_id")
        self.assertEqual(trace_stages, ["query_rewrite_for_embedding", "display_option_grouping", "range_selection", "quantity_determination", "final_explanation"])

    def test_query_validate_output_path_requires_overwrite_for_existing_output(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "query.xlsx"
            output_path.write_text("existing", encoding="utf-8")

            with self.assertRaisesRegex(ValueError, "输出已存在，请加 --overwrite"):
                query_estimate_llm.validate_output_path(output_path, overwrite=False)
            query_estimate_llm.validate_output_path(output_path, overwrite=True)

    def test_run_ingest_batch_parses_batch_id_from_filename(self):
        input_path = Path("excel_inputs/audit_ocr_export_20260630_001.xlsx")

        self.assertEqual(run_ingest_batch.infer_batch_id(input_path), "20260630_001")
        self.assertEqual(run_ingest_batch.batch_id_from_args(input_path, "manual_001"), "manual_001")

        with self.assertRaisesRegex(ValueError, "无法从文件名解析 batch_id"):
            run_ingest_batch.infer_batch_id(Path("excel_inputs/audit_ocr_export.xlsx"))

    def test_merge_samples_validate_output_paths_requires_overwrite_for_output_or_report(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)
            output_path = tmp_path / "samples" / "cost_item_samples_all.xlsx"
            report_path = merge_samples.dedup_report_path(output_path)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            report_path.write_text("existing", encoding="utf-8")

            with self.assertRaisesRegex(ValueError, "输出已存在，请加 --overwrite"):
                merge_samples.validate_output_paths(output_path, report_path, overwrite=False)
            merge_samples.validate_output_paths(output_path, report_path, overwrite=True)

    def test_legacy_query_entrypoint_removed_and_readme_points_to_official_script(self):
        legacy_name = "query_cost_" "item_estimate.py"
        self.assertFalse((ROOT / "estimator" / "scripts" / legacy_name).exists())
        readme = (ROOT / "README.md").read_text(encoding="utf-8")
        self.assertIn("estimator/scripts/query_cost_estimate_llm.py", readme)
        self.assertNotIn(f"estimator/scripts/{legacy_name}", readme)
        self.assertNotIn("project_name_embeddings.npy", readme)
        self.assertNotIn("project_detail_embeddings.npy", readme)
        self.assertNotIn("item_text_embeddings.npy", readme)


if __name__ == "__main__":
    unittest.main()
